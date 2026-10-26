"""Add and update people from a file: an Excel workbook (.xlsx, the first sheet) or a CSV file, one row per person.

Columns are found by their heading (first non-empty row), whatever the order and capitals: employee code (also
"employee number", "emp no", "staff id", "code", ...), name, and optionally department, shift and site (a site's
name). Other columns are ignored. Excel turns 1001 into a number; it is read back as the text "1001" (but an Excel
cell that dropped leading zeros, 0012 -> 12, can't be undone here: format that column as Text in Excel).

plan() checks every row against the database without changing anything (the preview):
- new: an employee code GateSense doesn't have yet;
- update: an existing person whose name, department, shift or site differ; only the columns the file has are
  compared, and a blank cell in such a column clears that field (the site keeps its value: every person has one);
- unchanged: an existing person the file matches;
- error: no code or name, a code that can't be a folder name, the same code twice in the file, a code that
  differs from an existing one only in capitals (they would share a photo folder on Windows), an unknown site.
People GateSense has but the file doesn't list are left alone. An inactive person is updated but not reactivated.

import_people() makes the plan again inside one write transaction and refuses the whole file if any row has an
error, so the file is imported completely or not at all. Every person added or changed gets its own audit row
(add_person / edit_person), plus one import_people row for the file.
"""

import csv
import io
import re
import sqlite3
from dataclasses import dataclass, field

from app import config, db, enrollment

FIELDS = ("employee_code", "name", "department", "shift", "site")
HEADINGS = {  # a heading, lower case with only letters and digits kept -> field
    "employee_code": ("employeecode", "employeenumber", "employeeno", "employeenum", "employeeid", "empcode",
                      "empno", "empnumber", "empid", "staffcode", "staffno", "staffnumber", "staffid", "code",
                      "number", "no", "id"),
    "name": ("name", "fullname", "employeename", "staffname", "personname"),
    "department": ("department", "dept", "departmentname", "section"),
    "shift": ("shift",),
    "site": ("site", "sitename", "location"),
}
TEMPLATE_HEADINGS = ("Employee code", "Name", "Department", "Shift", "Site")


class ImportFileError(ValueError):
    """The file as a whole can't be read (not Excel or CSV, no heading row, no code or name column, too many rows)."""


@dataclass
class ImportRow:
    line: int  # the row number in the file (1 = the heading row)
    employee_code: str
    name: str
    department: str | None = None
    shift: str | None = None
    site: str | None = None  # as written in the file
    site_id: int | None = None
    action: str = "new"  # new, update, unchanged, error
    person_id: int | None = None
    changes: dict[str, list] = field(default_factory=dict)  # field -> [old, new], for an update
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class ImportPlan:
    rows: list[ImportRow]
    columns: list[str]  # the fields found, in FIELDS order
    ignored_columns: list[str]  # headings that aren't one of the fields

    def count(self, action: str) -> int:
        return sum(r.action == action for r in self.rows)

    @property
    def errors(self) -> list[ImportRow]:
        return [r for r in self.rows if r.action == "error"]


def _key(heading) -> str:
    return re.sub(r"[^a-z0-9]", "", str(heading or "").lower())


def _cell(value) -> str:
    """A cell as text: Excel numbers 1001 / 1001.0 -> "1001", None -> ""."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(data: bytes, filename: str = "") -> list[list[str]]:
    """The file's rows as text cells: an .xlsx workbook's first sheet, or CSV (UTF-8, or Windows' ANSI codepage)."""
    if data[:2] == b"PK" or filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001  any reason the workbook can't be opened
            raise ImportFileError(f"this isn't an Excel workbook GateSense can read ({type(e).__name__}); "
                                  "save it as .xlsx or CSV") from None
        try:
            sheet = book.worksheets[0]
            return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or filename.lower().endswith(".xls"):
        raise ImportFileError("old Excel files (.xls) aren't supported: save it as .xlsx (Excel Workbook) or CSV")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("the file isn't text: save it as .xlsx or CSV (UTF-8)")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def parse(data: bytes, filename: str = "") -> tuple[list[ImportRow], list[str], list[str]]:
    """The file's people as rows (blank lines skipped), the fields found and the ignored headings."""
    if not data:
        raise ImportFileError("the file is empty")
    table = read_table(data, filename)
    start = next((i for i, row in enumerate(table) if any(row)), None)
    if start is None:
        raise ImportFileError("the file has no rows")
    headings = table[start]
    columns: dict[str, int] = {}
    ignored = []
    for index, heading in enumerate(headings):
        key = _key(heading)
        match = next((f for f, names in HEADINGS.items() if key in names), None)
        if match and match not in columns:
            columns[match] = index
        elif heading:
            ignored.append(heading)
    missing = [f for f in ("employee_code", "name") if f not in columns]
    if missing:
        raise ImportFileError(
            "the heading row needs an employee code column and a name column (e.g. \"Employee code\", \"Name\"); "
            f"found: {', '.join(h for h in headings if h) or 'nothing'}")
    rows = []
    for offset, cells in enumerate(table[start + 1:], start=start + 2):
        if not any(cells):
            continue
        value = {f: (cells[i] if i < len(cells) else "") for f, i in columns.items()}
        rows.append(ImportRow(line=offset, employee_code=value.get("employee_code", ""), name=value.get("name", ""),
                              department=value.get("department") or None, shift=value.get("shift") or None,
                              site=value.get("site") or None))
    if not rows:
        raise ImportFileError("the file has a heading row but no people under it")
    if len(rows) > config.IMPORT_MAX_ROWS:
        raise ImportFileError(f"{len(rows)} rows: at most {config.IMPORT_MAX_ROWS} people per file")
    return rows, [f for f in FIELDS if f in columns], ignored


def plan(conn: sqlite3.Connection, data: bytes, filename: str = "") -> ImportPlan:
    """What importing the file would do (see the module docstring); changes nothing."""
    rows, columns, ignored = parse(data, filename)
    people = {r["employee_code"]: dict(r) for r in conn.execute("SELECT * FROM persons")}
    by_folded = {code.lower(): code for code in people}
    sites = {r["name"].lower(): dict(r) for r in conn.execute("SELECT id, name FROM sites")}
    site_names = {s["id"]: s["name"] for s in sites.values()}
    seen: dict[str, int] = {}
    for row in rows:
        if not row.employee_code:
            row.errors.append("no employee code")
        else:
            try:
                row.employee_code = enrollment.validate_code(row.employee_code)
            except ValueError as e:
                row.errors.append(str(e))
            first = seen.setdefault(row.employee_code.lower(), row.line)
            if first != row.line:
                row.errors.append(f"the same employee code as line {first}")
        if not row.name:
            row.errors.append("no name")
        if row.site:
            site = sites.get(row.site.lower())
            if site is None:
                row.errors.append(f"no site called {row.site!r} (add it in Settings > Sites first)")
            else:
                row.site_id = site["id"]
        existing = people.get(row.employee_code)
        folded = by_folded.get(row.employee_code.lower())
        if existing is None and folded is not None:
            row.errors.append(f"differs only in capitals from the existing employee code {folded!r} (the two would "
                              "share one photo folder)")
        if row.errors:
            row.action = "error"
            continue
        if existing is None:
            row.action = "new"
            continue
        row.person_id = existing["id"]
        wanted = {"name": row.name}
        for name in ("department", "shift"):
            if name in columns:
                wanted[name] = getattr(row, name)
        if "site" in columns and row.site_id is not None:
            wanted["site_id"] = row.site_id
        for name, new in wanted.items():
            if (existing[name] or None) != (new or None):
                old = existing[name]
                if name == "site_id":
                    row.changes["site"] = [site_names.get(old), site_names.get(new)]
                else:
                    row.changes[name] = [old, new]
        row.action = "update" if row.changes else "unchanged"
        if not existing["active"]:
            row.notes.append("inactive: updated, but not reactivated")
    return ImportPlan(rows, columns, ignored)


def import_people(conn: sqlite3.Connection, data: bytes, filename: str, user: str, reason: str) -> ImportPlan:
    """Import the file in one transaction: every new person added, every changed one updated, all audited. Raises
    ValueError (nothing changed) if any row has an error."""
    db.require_reason(reason)
    conn.execute("BEGIN IMMEDIATE")  # planned and written under one write lock: nothing changes in between
    try:
        result = plan(conn, data, filename)
        if result.errors:
            raise ValueError("; ".join(f"line {r.line}: {', '.join(r.errors)}" for r in result.errors[:10])
                             + (f" (and {len(result.errors) - 10} more rows)" if len(result.errors) > 10 else ""))
        for row in result.rows:
            if row.action == "new":
                cur = conn.execute(
                    f"INSERT INTO persons (employee_code, name, department, shift, site_id) "
                    f"VALUES (?, ?, ?, ?, COALESCE(?, {db.DEFAULT_SITE}))",
                    (row.employee_code, row.name, row.department, row.shift, row.site_id))
                row.person_id = cur.lastrowid
                db.write_audit(conn, user, "add_person", f"person:{row.employee_code}",
                               {"name": row.name, "department": row.department, "shift": row.shift,
                                "site": row.site, "file": filename}, reason)
            elif row.action == "update":
                values = {("site_id" if k == "site" else k): v[1] for k, v in row.changes.items()}
                if "site_id" in values:
                    values["site_id"] = row.site_id
                sets = ", ".join(f"{k} = ?" for k in values)
                conn.execute(f"UPDATE persons SET {sets} WHERE id = ?", (*values.values(), row.person_id))
                db.write_audit(conn, user, "edit_person", f"person:{row.employee_code}",
                               {**row.changes, "file": filename}, reason)
        db.write_audit(conn, user, "import_people", f"file:{filename or 'upload'}",
                       {"added": result.count("new"), "updated": result.count("update"),
                        "unchanged": result.count("unchanged"), "rows": len(result.rows)}, reason)
        conn.commit()
    except BaseException:
        conn.rollback()
        raise
    return result


def template(kind: str = "xlsx") -> bytes:
    """An empty file with the headings and one example row, to fill in."""
    example = ("1001", "Ayesha Khan", "Stores", "Day", config.DEFAULT_SITE_NAME)
    if kind == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\r\n")
        writer.writerows([TEMPLATE_HEADINGS, example])
        return buf.getvalue().encode("utf-8-sig")
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "People"
    sheet.append(TEMPLATE_HEADINGS)
    sheet.append(example)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDE", (16, 28, 20, 12, 20)):
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["A"].number_format = "@"  # the codes as text: Excel keeps leading zeros typed there
    for cell in sheet["A"]:
        cell.number_format = "@"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
