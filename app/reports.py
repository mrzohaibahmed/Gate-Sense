"""Attendance reports for any date range: the daily register and the per-person totals, with CSV and Excel
export (openpyxl)."""

import io
import sqlite3
from datetime import date

from app import db
from app.shifts import (
    ABSENT, CALENDAR_CAME_IN, CAME_IN, CHECK, COMPLETE, EARLY_LEAVE, LATE, MISSING_IN, MISSING_OUT, NOT_IN_YET, OFF_DAY,
    ON_TIME, UNSCHEDULED,
)

# Both build on db.attendance_days, the same rule as /api/attendance, the overview and scripts/today.py: with shifts on,
# roster days (app/shifts.py: one day per roster row, plus unscheduled dates with passes; a night shift is dated the day
# it starts; "present" = on time, late, early leave, missing OUT); with shifts off, calendar days (first IN, last OUT;
# complete, missing IN, missing OUT, check; "present" = an IN that day), and the shift columns are left out.
# The Status column says "inside" for a day still running without an OUT (shifts off: today), like the screens.
INSIDE = "inside"

DAILY_COLUMNS_RANGE = [  # (header, key, kind) for exports
    ("Date", "date", "date"), ("Employee code", "employee_code", "text"), ("Name", "name", "text"),
    ("Department", "department", "text"), ("Shift", "shift_label", "text"), ("Shift start", "shift_start", "time"),
    ("Shift end", "shift_end", "time"), ("First IN", "first_in", "time"), ("Last OUT", "last_out", "time"),
    ("Hours worked", "worked_hours", "hours"), ("Late (min)", "late_minutes", "int"),
    ("Early leave (min)", "early_leave_minutes", "int"), ("Overtime (min)", "overtime_minutes", "int"),
    ("Status", "status_label", "text"),
]
CALENDAR_DAILY_COLUMNS = [
    ("Date", "date", "date"), ("Employee code", "employee_code", "text"), ("Name", "name", "text"),
    ("Department", "department", "text"), ("First IN", "first_in", "time"), ("Last OUT", "last_out", "time"),
    ("Hours worked", "worked_hours", "hours"), ("Status", "status_label", "text"),
]
CALENDAR_TOTAL_COLUMNS = [
    ("Employee code", "employee_code", "text"), ("Name", "name", "text"), ("Department", "department", "text"),
    ("Days", "days", "int"), ("Present", "days_present", "int"), ("Complete", "days_complete", "int"),
    ("Missing IN", "days_missing_in", "int"), ("Missing OUT", "days_missing_out", "int"),
    ("Days to check", "days_check", "int"), ("Total hours", "total_hours", "hours"),
    ("Average hours per worked day", "avg_hours", "hours"),
]
TOTAL_COLUMNS = [
    ("Employee code", "employee_code", "text"), ("Name", "name", "text"), ("Department", "department", "text"),
    ("Days", "days", "int"), ("Scheduled", "days_scheduled", "int"), ("Present", "days_present", "int"),
    ("On time", "days_on_time", "int"), ("Late", "days_late", "int"), ("Late (min)", "late_minutes", "int"),
    ("Early leave", "days_early_leave", "int"), ("Early leave (min)", "early_leave_minutes", "int"),
    ("Overtime (min)", "overtime_minutes", "int"), ("Missing IN", "days_missing_in", "int"),
    ("Missing OUT", "days_missing_out", "int"), ("Days to check", "days_check", "int"),
    ("Absent", "days_absent", "int"), ("Not in yet", "days_not_in_yet", "int"), ("Off days", "days_off", "int"),
    ("Unscheduled", "days_unscheduled", "int"), ("Total hours", "total_hours", "hours"),
    ("Average hours per worked day", "avg_hours", "hours"),
]
COUNTED = {  # PersonTotal field -> the status it counts
    "days_on_time": ON_TIME, "days_late": LATE, "days_early_leave": EARLY_LEAVE, "days_missing_in": MISSING_IN,
    "days_missing_out": MISSING_OUT, "days_check": CHECK, "days_absent": ABSENT, "days_not_in_yet": NOT_IN_YET,
    "days_off": OFF_DAY, "days_unscheduled": UNSCHEDULED, "days_complete": COMPLETE,
}
MINUTES = ("late_minutes", "early_leave_minutes", "overtime_minutes")


def columns(shifts: bool) -> tuple[list, list]:
    """(daily register columns, per-person total columns) for exports: with shifts on or off."""
    return (DAILY_COLUMNS_RANGE, TOTAL_COLUMNS) if shifts else (CALENDAR_DAILY_COLUMNS, CALENDAR_TOTAL_COLUMNS)


def range_rows(conn: sqlite3.Connection, first: date, last: date, department: str | None = None,
               site_id: int | None = None) -> list[dict]:
    """One row per day from `first` to `last` (db.attendance_days), by date then name. Read-only."""
    people = {r["employee_code"]: dict(r) for r in db.list_persons(conn)}
    rows = []
    for d in db.attendance_days(conn, first, last):
        p = people[d.person]
        if department and (p["department"] or "") != department:
            continue
        if site_id is not None and p["site_id"] != site_id:
            continue
        rows.append({"date": d.day, "employee_code": d.person, "name": p["name"], "department": p["department"],
                     "shift": d.shift, "shift_label": d.shift or ("OFF" if d.status == OFF_DAY else ""),
                     "shift_start": d.start, "shift_end": d.end, "first_in": d.first_in, "last_out": d.last_out,
                     "worked_hours": None if d.worked_minutes is None else round(d.worked_minutes / 60, 2),
                     "late_minutes": d.late_minutes, "early_leave_minutes": d.early_leave_minutes,
                     "overtime_minutes": d.overtime_minutes, "status": d.status, "open": d.open,
                     "status_label": INSIDE if d.status == MISSING_OUT and d.open else d.status})
    return sorted(rows, key=lambda r: (r["date"], r["name"].lower(), r["employee_code"]))


def person_totals(rows: list[dict], shifts: bool = True) -> list[dict]:
    """Per person, from the daily rows only (so the totals always add up to them): days in the register, scheduled
    days (a shift, not OFF), days present (came in for the shift), days per status, late / early leave / overtime
    minutes, total hours and the average over the days with hours."""
    totals: dict[str, dict] = {}
    for r in rows:
        t = totals.setdefault(r["employee_code"], {
            "employee_code": r["employee_code"], "name": r["name"], "department": r["department"], "days": 0,
            "days_scheduled": 0, "days_present": 0, **{k: 0 for k in COUNTED}, **{k: 0 for k in MINUTES},
            "total_hours": 0.0, "avg_hours": None, "_worked_days": 0})
        t["days"] += 1
        t["days_scheduled"] += r["shift"] is not None
        t["days_present"] += r["status"] in (CAME_IN if shifts else CALENDAR_CAME_IN)
        for key, status in COUNTED.items():
            t[key] += r["status"] == status and not (status == MISSING_OUT and r["open"])  # open: still inside
        for key in MINUTES:
            t[key] += r[key] or 0
        if r["worked_hours"] is not None:
            t["total_hours"] += r["worked_hours"]
            t["_worked_days"] += 1
    out = []
    for t in totals.values():
        worked = t.pop("_worked_days")
        t["total_hours"] = round(t["total_hours"], 2)
        t["avg_hours"] = round(t["total_hours"] / worked, 2) if worked else None
        out.append(t)
    return sorted(out, key=lambda t: (t["name"].lower(), t["employee_code"]))


def _cell(value, kind: str):
    if value is None:
        return ""
    if kind == "time":
        return value.strftime("%H:%M:%S")
    if kind == "date":
        return value.isoformat()
    if kind == "yesno":
        return "YES" if value else ""
    return value


def rows_csv(columns: list[tuple[str, str, str]], rows: list[dict]) -> bytes:
    """CSV with a BOM so Excel opens it as UTF-8."""
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([h for h, _, _ in columns])
    for r in rows:
        writer.writerow([_cell(r.get(k), kind) for _, k, kind in columns])
    return buf.getvalue().encode("utf-8-sig")


def rows_xlsx(sheets: list[tuple[str, list[tuple[str, str, str]], list[dict]]], title: str) -> bytes:
    """One sheet per (name, columns, rows): bold white header on the GateSense teal, first row frozen, filters,
    real dates, times (hh:mm) and numbers, column widths fitted. `title` goes in the document properties."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = title
    wb.properties.creator = "GateSense"
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0E7C66")
    for name, columns, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append([h for h, _, _ in columns])
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            values = []
            for _, key, kind in columns:
                v = r.get(key)
                values.append(v.time() if kind == "time" and v is not None else
                              ("YES" if v else "") if kind == "yesno" else ("" if v is None else v))
            ws.append(values)
        formats = {"date": "yyyy-mm-dd", "time": "hh:mm", "hours": "0.00", "int": "0"}
        for i, (header, _, kind) in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            if kind in formats:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = formats[kind]
            width = max([len(header)] + [len(str(_cell(r.get(columns[i - 1][1]), kind))) for r in rows[:500]])
            ws.column_dimensions[letter].width = min(width + 2, 40)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
