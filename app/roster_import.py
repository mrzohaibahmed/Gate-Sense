"""The roster from a file: an Excel workbook (.xlsx, the first sheet) or a CSV file, one row per person and one
column per date, like the Roster screen.

Headings (the first non-empty row): an employee code column (the same names People > Import accepts: "Employee code",
"Emp no", "Staff ID", "Code", ...), optionally a name column (shown in the preview, never imported), and one column
per date written as YYYY-MM-DD ("Mon 2026-10-26" and Excel date cells work too). Other headings are ignored. A cell is
a shift's name (any capitals), OFF, or blank: a blank cell leaves that day as it is (remove an entry on the screen).

plan() checks every cell against the database without changing anything (the preview): a cell is a change (set) or
unchanged; errors, each with its line and column: no or unknown employee code, an inactive person, the same code
twice, an unknown shift. A file GateSense can't use at all (not Excel or CSV, no code column, no date columns, a date
heading that isn't YYYY-MM-DD, the same date twice, too many dates or rows) raises ImportFileError.

import_roster() plans again inside one write transaction and refuses the whole file while any cell has an error;
otherwise every change is written like a roster edit on the screen (a set_roster audit row each) plus one
import_roster row for the file.
"""

import csv
import io
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime

from app import config, db
from app.people_import import HEADINGS, ImportFileError, _key, read_table
from app.shifts import OFF

ISO_DATE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
OTHER_DATE = re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b|\b\d{4}/\d{1,2}/\d{1,2}\b")


@dataclass
class CellPlan:
    line: int  # the row number in the file (1 = the heading row)
    column: str  # the date heading as written
    employee_code: str
    name: str  # the person's name in GateSense
    person_id: int
    date: date
    old: str | None  # the roster entry now: a shift's name, OFF, or None (no entry)
    new: str  # a shift's name or OFF
    shift_id: int | None  # None = OFF
    action: str = "set"  # set or unchanged


@dataclass
class CellError:
    line: int
    column: str  # the heading of the cell ("" for a problem with the whole row)
    employee_code: str
    message: str


@dataclass
class RosterPlan:
    dates: list[date]
    cells: list[CellPlan] = field(default_factory=list)
    errors: list[CellError] = field(default_factory=list)
    people: int = 0  # rows with a known person
    ignored_columns: list[str] = field(default_factory=list)

    def count(self, action: str) -> int:
        return sum(c.action == action for c in self.cells)


@dataclass
class _Row:
    line: int
    code: str
    cells: dict[int, str]  # column index -> cell text


def parse(data: bytes, filename: str = "") -> tuple[dict[int, date], list[_Row], list[str], dict[int, str]]:
    """The date columns (index -> date), the rows, the ignored headings and every heading by index."""
    if not data:
        raise ImportFileError("the file is empty")
    table = read_table(data, filename)
    start = next((i for i, row in enumerate(table) if any(row)), None)
    if start is None:
        raise ImportFileError("the file has no rows")
    headings = dict(enumerate(table[start]))
    code_col = name_col = None
    dates: dict[int, date] = {}
    ignored = []
    for index, heading in headings.items():
        key = _key(heading)
        found = ISO_DATE.search(heading)
        if found:
            try:
                day = date(*map(int, found.groups()))
            except ValueError:
                raise ImportFileError(f"the heading {heading!r} isn't a real date") from None
            if day in dates.values():
                raise ImportFileError(f"{day.isoformat()} is in two columns")
            dates[index] = day
        elif OTHER_DATE.search(heading):
            raise ImportFileError(f"write the dates in the heading row as YYYY-MM-DD (e.g. 2026-10-26), not "
                                  f"{heading!r}: day and month order can't be guessed")
        elif code_col is None and key in HEADINGS["employee_code"]:
            code_col = index
        elif name_col is None and key in HEADINGS["name"]:
            name_col = index
        elif heading:
            ignored.append(heading)
    if code_col is None:
        raise ImportFileError("the heading row needs an employee code column (e.g. \"Employee code\"); found: "
                              + (", ".join(h for h in headings.values() if h) or "nothing"))
    if not dates:
        raise ImportFileError("the heading row has no dates: one column per day, headed YYYY-MM-DD")
    if len(dates) > config.ROSTER_MAX_DAYS:
        raise ImportFileError(f"{len(dates)} dates: at most {config.ROSTER_MAX_DAYS} per file")
    rows = []
    for line, cells in enumerate(table[start + 1:], start=start + 2):
        if not any(cells):
            continue
        rows.append(_Row(line, cells[code_col] if code_col < len(cells) else "",
                         {i: cells[i] for i in dates if i < len(cells) and cells[i]}))
    if not rows:
        raise ImportFileError("the file has a heading row but no people under it")
    if len(rows) > config.IMPORT_MAX_ROWS:
        raise ImportFileError(f"{len(rows)} rows: at most {config.IMPORT_MAX_ROWS} people per file")
    return dates, rows, ignored, headings


def plan(conn: sqlite3.Connection, data: bytes, filename: str = "") -> RosterPlan:
    """What importing the file would do (see the module docstring); changes nothing."""
    dates, rows, ignored, headings = parse(data, filename)
    result = RosterPlan(sorted(dates.values()), ignored_columns=ignored)
    people = {r["employee_code"]: dict(r) for r in conn.execute("SELECT id, employee_code, name, active FROM persons")}
    shifts = {s.name.lower(): s for s in db.list_shifts(conn)}
    current = {(r["person_id"], r["date"]): r["shift_id"]
               for r in db.roster_rows(conn, min(dates.values()), max(dates.values()))}
    names = {s.id: s.name for s in shifts.values()}
    seen: dict[str, int] = {}
    for row in rows:
        person = people.get(row.code)
        problem = None
        if not row.code:
            problem = "no employee code"
        elif person is None:
            problem = f"no person with employee code {row.code!r} (add them in People first)"
        elif not person["active"]:
            problem = f"{person['name']} is inactive (activate them in People first)"
        elif seen.setdefault(row.code, row.line) != row.line:
            problem = f"the same employee code as line {seen[row.code]}"
        if problem:
            result.errors.append(CellError(row.line, "", row.code, problem))
            continue
        result.people += 1
        for index, text in sorted(row.cells.items()):
            day, value = dates[index], text.strip()
            if value.upper() == OFF:
                shift_id, new = None, OFF
            elif value.lower() in shifts:
                shift = shifts[value.lower()]
                shift_id, new = shift.id, shift.name
            else:
                result.errors.append(CellError(row.line, headings[index], row.code,
                                               f"no shift called {value!r} (use a shift's name, or OFF)"))
                continue
            key = (person["id"], day.isoformat())
            old = None if key not in current else OFF if current[key] is None else names.get(current[key])
            result.cells.append(CellPlan(row.line, headings[index], row.code, person["name"], person["id"], day,
                                         old, new, shift_id, "unchanged" if old == new else "set"))
    return result


def import_roster(conn: sqlite3.Connection, data: bytes, filename: str, user: str, reason: str) -> RosterPlan:
    """Import the file in one transaction (see the module docstring). Raises ValueError (nothing changed) if any cell
    has an error."""
    db.require_reason(reason)
    conn.execute("BEGIN IMMEDIATE")  # planned and written under one write lock: nothing changes in between
    try:
        result = plan(conn, data, filename)
        if result.errors:
            shown = "; ".join(f"line {e.line}{' ' + e.column if e.column else ''}: {e.message}"
                              for e in result.errors[:10])
            raise ValueError(shown + (f" (and {len(result.errors) - 10} more)" if len(result.errors) > 10 else ""))
        changes = [(c.person_id, c.date, c.shift_id) for c in result.cells if c.action == "set"]
        db.write_roster(conn, changes, reason, user)
        db.write_audit(conn, user, "import_roster", f"file:{filename or 'upload'}",
                       {"changed": len(changes), "unchanged": result.count("unchanged"), "people": result.people,
                        "from": result.dates[0].isoformat(), "to": result.dates[-1].isoformat()}, reason)
        conn.commit()
    except BaseException:
        conn.rollback()
        raise
    return result


def template(conn: sqlite3.Connection, first: date, last: date, kind: str = "xlsx", site_id: int | None = None,
             department: str | None = None) -> bytes:
    """The roster of `first` to `last` as a file to fill in: every active person (of the site / department), one
    column per date, each cell the current entry (a shift's name, OFF, or blank). Excel also gets a Shifts sheet."""
    from datetime import timedelta

    days = [first + timedelta(days=i) for i in range((last - first).days + 1)]
    people = [p for p in db.list_persons(conn) if p["active"] and (site_id is None or p["site_id"] == site_id)
              and (not department or (p["department"] or "") == department)]
    shifts = db.list_shifts(conn, first)
    names = {s.id: s.name for s in shifts}
    entries = {(r["person_id"], r["date"]): r["shift_id"] for r in db.roster_rows(conn, first, last)}

    def cell(person_id: int, day: date) -> str:
        key = (person_id, day.isoformat())
        return "" if key not in entries else OFF if entries[key] is None else names.get(entries[key], "")

    rows = [[p["employee_code"], p["name"], *(cell(p["id"], d) for d in days)] for p in people]
    if kind == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\r\n")
        writer.writerow(["Employee code", "Name", *(d.isoformat() for d in days)])
        writer.writerows(rows)
        return buf.getvalue().encode("utf-8-sig")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Roster"
    sheet.append(["Employee code", "Name", *(datetime.combine(d, datetime.min.time()) for d in days)])
    for values in rows:
        sheet.append(values)
    for cell_ in sheet[1]:
        cell_.font = Font(bold=True)
    for index in range(3, len(days) + 3):
        sheet.cell(row=1, column=index).number_format = "ddd yyyy-mm-dd"
        sheet.column_dimensions[get_column_letter(index)].width = 16
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 28
    for (code_cell,) in sheet.iter_rows(min_col=1, max_col=1):
        code_cell.number_format = "@"  # codes as text: Excel keeps leading zeros typed there
    sheet.freeze_panes = "C2"
    help_sheet = book.create_sheet("Shifts")
    help_sheet.append(["Write in a day's cell", "Times"])
    for s in shifts:
        help_sheet.append([s.name, f"{s.start:%H:%M}-{s.end:%H:%M}" + (" (next day)" if s.crosses_midnight else "")])
    help_sheet.append([OFF, "day off"])
    help_sheet.append(["(blank)", "leaves that day as it is"])
    for cell_ in help_sheet[1]:
        cell_.font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 22
    help_sheet.column_dimensions["B"].width = 26
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
