"""Shifts and roster endpoints (api/routes/shifts.py, api/routes/roster.py): who may do what, effective-dated shift
changes, the roster grid and its saves, and the roster file import (template, preview with every error, all or
nothing)."""

import csv
import io
from datetime import date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import create_app
from app import db
from tests.test_data_api import as_role, first_password, fixed_clock, items  # noqa: F401  (autouse fixtures)

MON = date(2026, 10, 19)
WEEK = [MON + timedelta(days=i) for i in range(7)]


@pytest.fixture
def app():
    application = create_app()
    with TestClient(application):
        conn = db.connect()
        db.sync_persons(conn, ["1001", "1002", "1003"])
        for code, name, dept in (("1001", "Ayesha", "Stores"), ("1002", "Bilal", "Office"),
                                 ("1003", "Chand", "Stores")):
            conn.execute("UPDATE persons SET name = ?, department = ? WHERE employee_code = ?", (name, dept, code))
        conn.commit()
        conn.close()
        yield application


def ids(client: TestClient) -> tuple[dict, dict]:
    """(employee code -> person id, shift name -> shift id)"""
    people = {p["employee_code"]: p["id"] for p in items(client.get("/api/people"))}
    shifts = {s["name"]: s["id"] for s in items(client.get("/api/shifts"))}
    return people, shifts


def grid(client: TestClient, **params) -> dict[str, dict[str, int | None]]:
    """employee code -> {date: shift id or None (OFF)} for the week of MON."""
    rows = items(client.get("/api/roster", params={"from": MON.isoformat(), **params}))
    return {r["employee_code"]: {c["date"]: c["shift_id"] for c in r["cells"]} for r in rows}


def audit(app, action: str) -> list[dict]:
    return [e for e in app.state.db.audit_entries() if e["action"] == action]


# --- Who may do what ------------------------------------------------------------------------------------------


def test_admin_and_hr_edit_the_roster_viewer_reads_it_guard_has_no_access(app):
    for role in ("admin", "hr"):
        c = as_role(app, role)
        people, shifts = ids(c)
        r = c.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(),
                                                    "shift_id": shifts["Morning"]}], "reason": f"{role} plans"})
        assert r.status_code == 200, (role, r.text)
        assert c.patch(f"/api/shifts/{shifts['Night']}", json={"grace_min": 20, "reason": "x"}).status_code == 200
    viewer = as_role(app, "viewer")
    assert grid(viewer)["1001"] == {MON.isoformat(): ids(viewer)[1]["Morning"]}
    assert viewer.put("/api/roster", json={"entries": [{"person_id": 1, "date": MON.isoformat()}],
                                           "reason": "x"}).status_code == 403
    assert viewer.post("/api/shifts", json={"name": "Late", "start": "14:00", "end": "22:00",
                                            "reason": "x"}).status_code == 403
    assert viewer.get("/api/roster/import/template").status_code == 403
    guard = as_role(app, "guard")
    for path in ("/api/roster", "/api/shifts", "/api/shifts/1"):
        assert guard.get(path).status_code == 403, path
    denied = " ".join(a["details"] for a in audit(app, "permission_denied"))
    assert all(p in denied for p in ("roster:manage", "roster:read", "shifts:manage"))


# --- Shifts ------------------------------------------------------------------------------------------------------


def test_shift_list_create_and_validation(app):
    admin = as_role(app, "admin")
    shifts = items(admin.get("/api/shifts"))
    assert [(s["name"], s["start"], s["end"], s["crosses_midnight"], s["effective_from"]) for s in shifts] == [
        ("Morning", "06:30", "14:30", False, None), ("General", "08:00", "16:30", False, None),
        ("Night", "22:00", "06:00", True, None)]
    assert shifts[0]["versions"][0]["effective_from"] is None and shifts[0]["roster_days"] == 0
    r = admin.post("/api/shifts", json={"name": "Late", "start": "14:00", "end": "22:00", "reason": "new shift"})
    assert r.status_code == 201 and (r.json()["grace_min"], r.json()["crosses_midnight"]) == (30, False)
    overnight = admin.post("/api/shifts", json={"name": "Graveyard", "start": "23:00", "end": "07:00",
                                                "grace_min": 10, "reason": "new"})
    assert overnight.json()["crosses_midnight"] is True  # worked out from the times
    assert admin.post("/api/shifts", json={"name": "late", "start": "14:00", "end": "22:00",
                                           "reason": "x"}).status_code == 409
    bad = admin.post("/api/shifts", json={"name": "Odd", "start": "22:00", "end": "06:00", "crosses_midnight": False,
                                          "reason": "x"})
    assert bad.status_code == 422 and "crosses midnight" in bad.text
    assert admin.post("/api/shifts", json={"name": "OFF", "start": "08:00", "end": "16:00",
                                           "reason": "x"}).status_code == 422
    assert admin.post("/api/shifts", json={"name": "Long", "start": "08:00", "end": "16:00", "grace_min": 900,
                                           "reason": "x"}).status_code == 422
    assert [a["details"] for a in audit(app, "add_shift")][-1].startswith('{"name": "Late"')


def test_a_shift_change_applies_from_its_date(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    morning = shifts["Morning"]
    wed = MON + timedelta(days=2)
    r = hr.patch(f"/api/shifts/{morning}", json={"start": "07:00", "end": "15:00", "effective_from": wed.isoformat(),
                                                 "reason": "new hours"})
    assert r.status_code == 200
    assert [(v["effective_from"], v["start"], v["planned"]) for v in r.json()["versions"]] == [
        (None, "06:30", False), (wed.isoformat(), "07:00", False)]
    before = hr.get(f"/api/shifts/{morning}", params={"on": MON.isoformat()}).json()
    assert (before["start"], before["effective_from"]) == ("06:30", None)
    assert hr.get(f"/api/shifts/{morning}").json()["start"] == "07:00"  # today is after Wednesday the 21st
    entry = audit(app, "edit_shift")[0]
    assert '"effective_from": "2026-10-21"' in entry["details"] and entry["reason"] == "new hours"
    # Rename: applies everywhere at once.
    renamed = hr.patch(f"/api/shifts/{morning}", json={"name": "Early", "reason": "clearer"}).json()
    assert renamed["name"] == "Early" and len(renamed["versions"]) == 2
    assert hr.patch(f"/api/shifts/{morning}", json={"name": "general", "reason": "clash"}).status_code == 409
    # A planned change (after today) can be withdrawn; one in force can't.
    later = (date.today() + timedelta(days=10)).isoformat()
    planned = hr.patch(f"/api/shifts/{morning}", json={"grace_min": 15, "effective_from": later,
                                                       "reason": "next month"}).json()
    plan = next(v for v in planned["versions"] if v["effective_from"] == later)
    assert plan["planned"] and planned["grace_min"] == 30  # not in force yet
    in_force = next(v for v in planned["versions"] if v["effective_from"] == wed.isoformat())
    assert hr.delete(f"/api/shifts/{morning}/versions/{in_force['id']}",
                     params={"reason": "undo"}).status_code == 409
    withdrawn = hr.delete(f"/api/shifts/{morning}/versions/{plan['id']}", params={"reason": "plan dropped"})
    assert withdrawn.status_code == 200 and len(withdrawn.json()["versions"]) == 2
    assert audit(app, "delete_shift_version")[0]["reason"] == "plan dropped"


def test_a_rostered_shift_cant_be_deleted(app):
    admin = as_role(app, "admin")
    people, shifts = ids(admin)
    late = admin.post("/api/shifts", json={"name": "Late", "start": "14:00", "end": "22:00", "reason": "x"}).json()
    admin.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(),
                                                "shift_id": late["id"]}], "reason": "plan"})
    assert admin.get(f"/api/shifts/{late['id']}").json()["roster_days"] == 1
    r = admin.delete(f"/api/shifts/{late['id']}", params={"reason": "not needed"})
    assert r.status_code == 409 and "roster" in r.text
    admin.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(), "clear": True}],
                                   "reason": "replan"})
    assert admin.delete(f"/api/shifts/{late['id']}", params={"reason": "not needed"}).status_code == 204
    assert admin.get(f"/api/shifts/{late['id']}").status_code == 404


# --- Roster grid ---------------------------------------------------------------------------------------------------


def test_roster_grid_set_off_clear_and_bulk(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    entries = [{"person_id": people[code], "date": d.isoformat(), "shift_id": shifts["General"]}
               for code in ("1001", "1003") for d in WEEK[:5]]  # bulk: two people, Monday to Friday
    entries += [{"person_id": people["1001"], "date": d.isoformat(), "shift_id": None} for d in WEEK[5:]]  # OFF
    r = hr.put("/api/roster", json={"entries": entries, "reason": "week plan"})
    assert r.json() == {"changed": 12, "removed": 0}
    g = grid(hr)
    assert g["1001"][WEEK[0].isoformat()] == shifts["General"] and g["1001"][WEEK[6].isoformat()] is None
    assert g["1002"] == {} and len(g["1003"]) == 5
    assert hr.put("/api/roster", json={"entries": entries, "reason": "again"}).json() == {"changed": 0, "removed": 0}
    r = hr.put("/api/roster", json={"entries": [
        {"person_id": people["1003"], "date": WEEK[4].isoformat(), "clear": True},
        {"person_id": people["1003"], "date": WEEK[3].isoformat(), "shift_id": shifts["Night"]}], "reason": "swap"})
    assert r.json() == {"changed": 1, "removed": 1}
    assert grid(hr)["1003"][WEEK[3].isoformat()] == shifts["Night"] and WEEK[4].isoformat() not in grid(hr)["1003"]
    assert [len(audit(app, a)) for a in ("set_roster", "clear_roster")] == [13, 1]
    # Filters, paging and range
    assert set(grid(hr, department="Stores")) == {"1001", "1003"}
    assert grid(hr, site_id=99) == {}
    week = hr.get("/api/roster", params={"from": MON.isoformat()}).json()
    assert week["total"] == 3 and [r["name"] for r in week["items"]] == ["Ayesha", "Bilal", "Chand"]
    assert hr.get("/api/roster", params={"from": MON.isoformat(),
                                         "to": (MON + timedelta(days=62)).isoformat()}).status_code == 422
    default = hr.get("/api/roster").json()["items"][0]["cells"]
    monday = date.today() - timedelta(days=date.today().weekday())
    assert all(monday <= date.fromisoformat(c["date"]) <= monday + timedelta(days=6) for c in default)


def test_roster_saves_are_checked(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    one = {"person_id": people["1001"], "date": MON.isoformat(), "shift_id": shifts["General"]}
    assert hr.put("/api/roster", json={"entries": [one], "reason": " "}).status_code == 422
    assert hr.put("/api/roster", json={"entries": [one, one], "reason": "x"}).status_code == 422
    assert hr.put("/api/roster", json={"entries": [{**one, "person_id": 999}], "reason": "x"}).status_code == 404
    assert hr.put("/api/roster", json={"entries": [{**one, "shift_id": 999}], "reason": "x"}).status_code == 422
    assert hr.put("/api/roster", json={"entries": [], "reason": "x"}).status_code == 422
    assert grid(hr)["1001"] == {}  # nothing was written


def test_a_deactivated_person_with_entries_stays_on_the_grid(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    hr.put("/api/roster", json={"entries": [{"person_id": people["1002"], "date": MON.isoformat(),
                                             "shift_id": shifts["General"]}], "reason": "plan"})
    assert hr.post(f"/api/people/{people['1002']}/deactivate", json={"reason": "left"}).status_code == 200
    rows = {r["employee_code"]: r for r in items(hr.get("/api/roster", params={"from": MON.isoformat()}))}
    assert rows["1002"]["active"] is False and rows["1002"]["cells"]
    nxt = items(hr.get("/api/roster", params={"from": (MON + timedelta(days=7)).isoformat()}))
    assert "1002" not in {r["employee_code"] for r in nxt}


def test_roster_changes_show_on_the_attendance_at_once(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    conn = db.connect()
    cin, cout = db.upsert_camera(conn, "entrance", 0, "IN"), db.upsert_camera(conn, "exit", 1, "OUT")
    db.add_event(conn, people["1001"], cin, datetime(2026, 10, 19, 8, 5), 0.8, None)
    db.add_event(conn, people["1001"], cout, datetime(2026, 10, 19, 16, 35), 0.8, None)
    conn.close()
    hr.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(),
                                             "shift_id": shifts["Morning"]}], "reason": "plan"})
    day = {"from": MON.isoformat(), "to": MON.isoformat(), "person": people["1001"]}
    assert [(r["shift"], r["status"], r["late_minutes"]) for r in items(hr.get("/api/attendance", params=day))] == [
        ("Morning", "late", 95)]
    hr.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(),
                                             "shift_id": shifts["General"]}], "reason": "moved"})
    assert [(r["shift"], r["status"], r["late_minutes"]) for r in items(hr.get("/api/attendance", params=day))] == [
        ("General", "on time", 0)]


# --- Roster file import --------------------------------------------------------------------------------------------


def upload(client: TestClient, path: str, data: bytes, filename: str, **params):
    kind = "text/csv" if filename.endswith(".csv") else \
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return client.post(path, content=data, params={"filename": filename, **params}, headers={"Content-Type": kind})


def csv_file(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def test_the_template_round_trips_unchanged(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    hr.put("/api/roster", json={"entries": [{"person_id": people["1001"], "date": MON.isoformat(),
                                             "shift_id": shifts["Night"]},
                                            {"person_id": people["1001"], "date": WEEK[6].isoformat(),
                                             "shift_id": None}], "reason": "plan"})
    x = hr.get("/api/roster/import/template", params={"from": MON.isoformat()})
    assert x.status_code == 200 and 'gatesense-roster-2026-10-19_2026-10-25.xlsx' in x.headers["content-disposition"]
    book = load_workbook(io.BytesIO(x.content))
    assert book.sheetnames == ["Roster", "Shifts"]
    sheet = book["Roster"]
    assert [c.value for c in sheet[1]][:2] == ["Employee code", "Name"] and sheet.cell(1, 3).value.date() == MON
    assert [c.value for c in sheet[2]][:3] == ["1001", "Ayesha", "Night"] and sheet.cell(2, 9).value == "OFF"
    preview = upload(hr, "/api/roster/import/preview", x.content, "roster.xlsx").json()
    assert preview["counts"] == {"set": 0, "unchanged": 2, "error": 0} and preview["people"] == 3
    assert preview["dates"][0] == MON.isoformat() and len(preview["dates"]) == 7
    text = hr.get("/api/roster/import/template", params={"from": MON.isoformat(), "format": "csv",
                                                         "department": "Office"})
    lines = list(csv.reader(io.StringIO(text.content.decode("utf-8-sig"))))
    assert lines[0][:3] == ["Employee code", "Name", "2026-10-19"] and [r[0] for r in lines[1:]] == ["1002"]


def test_import_preview_lists_every_error_and_import_is_all_or_nothing(app):
    hr = as_role(app, "hr")
    people, shifts = ids(hr)
    hr.post(f"/api/people/{people['1003']}/deactivate", json={"reason": "left"})
    bad = csv_file([["Emp no", "Name", "Notes", "2026-10-19", "Tue 2026-10-20"],
                    ["1001", "Ayesha", "x", "morning", "OFF"],
                    ["1002", "Bilal", "", "Lunch", ""],
                    ["9999", "Nobody", "", "General", ""],
                    ["1003", "Chand", "", "General", ""],
                    ["1001", "Ayesha again", "", "Night", ""]])
    preview = upload(hr, "/api/roster/import/preview", bad, "week.csv").json()
    assert preview["counts"] == {"set": 2, "unchanged": 0, "error": 4}
    assert preview["ignored_columns"] == ["Notes"]
    assert [(c["employee_code"], c["date"], c["old"], c["new"]) for c in preview["changes"]] == [
        ("1001", "2026-10-19", None, "Morning"), ("1001", "2026-10-20", None, "OFF")]
    assert [(e["line"], e["column"], e["message"].split(" (")[0]) for e in preview["errors"]] == [
        (3, "2026-10-19", "no shift called 'Lunch'"), (4, "", "no person with employee code '9999'"),
        (5, "", "Chand is inactive"), (6, "", "the same employee code as line 2")]
    r = upload(hr, "/api/roster/import", bad, "week.csv", reason="week plan")
    assert r.status_code == 422 and "Nothing was imported" in r.text
    assert grid(hr)["1001"] == {} and not audit(app, "import_roster")

    good = csv_file([["Employee code", "2026-10-19", "2026-10-20", "2026-10-21"],
                     ["1001", "Morning", "OFF", ""], ["1002", "general", "", "Night"]])
    assert upload(hr, "/api/roster/import", good, "week.csv", reason=" ").status_code == 422
    r = upload(hr, "/api/roster/import", good, "week.csv", reason="week plan")
    assert r.json() == {"file": "week.csv", "changed": 4, "unchanged": 0}
    g = grid(hr)
    assert g["1001"] == {"2026-10-19": shifts["Morning"], "2026-10-20": None}  # the blank cell left alone
    assert g["1002"] == {"2026-10-19": shifts["General"], "2026-10-21": shifts["Night"]}
    assert len(audit(app, "set_roster")) == 4
    [entry] = audit(app, "import_roster")
    assert entry["reason"] == "week plan" and '"changed": 4' in entry["details"]


def test_roster_files_that_cant_be_used(app):
    hr = as_role(app, "hr")
    cases = {
        "no dates": [["Employee code", "Name"], ["1001", "Ayesha"]],
        "day/month": [["Employee code", "19/10/2026"], ["1001", "General"]],
        "no code column": [["Name", "2026-10-19"], ["Ayesha", "General"]],
        "a date twice": [["Employee code", "2026-10-19", "Mon 2026-10-19"], ["1001", "General", "OFF"]],
    }
    for name, rows in cases.items():
        r = upload(hr, "/api/roster/import/preview", csv_file(rows), "x.csv")
        assert r.status_code == 422, name
    assert "YYYY-MM-DD" in upload(hr, "/api/roster/import/preview", csv_file(cases["day/month"]), "x.csv").text
    assert upload(hr, "/api/roster/import/preview", b"", "x.csv").status_code == 422
