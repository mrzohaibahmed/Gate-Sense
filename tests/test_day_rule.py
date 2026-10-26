"""One attendance per person per roster day (app/shifts.py), proven end to end on the audit's cases a-h and the newer
ones (off day, unscheduled, night shift, roster change): the roster goes in through PUT /api/roster, camera passes
through the worker's Recorder (cool-down per person per direction, the database check), corrections through the API,
and every screen's data, the reports, the CSV/XLSX exports and scripts/today.py are read back and must agree.

The General shift is 08:00-16:30 (IN window 06:00-12:15, OUT window 12:15-20:30); Night is 22:00-06:00, dated the day
it starts. The clock is fixed at 18:00 today (tests/test_data_api.py fixed_clock), so today's General shift is still
running."""

import contextlib
import csv
import io
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import create_app
from app import db
from app.gallery import Match
from app.tracker import Sample, TrackState
from tests.test_data_api import as_role, first_password, fixed_clock  # noqa: F401  (autouse fixtures)

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import recognize  # noqa: E402
import today as today_script  # noqa: E402

D = date(2026, 10, 20)
TODAY = date.today()
CODE = "1001"
CAMS = {"IN": {"name": "entrance", "direction": "IN"}, "OUT": {"name": "exit", "direction": "OUT"},
        "IN2": {"name": "entrance2", "direction": "IN"}}
OVERVIEW = ("people", "present", "inside_now", "on_time", "late", "early_leave", "overtime", "missing_in",
            "missing_out", "check", "absent", "not_in_yet", "off_day", "unscheduled")
MONTHLY = ("days", "days_scheduled", "days_present", "days_late", "late_minutes", "days_early_leave",
           "early_leave_minutes", "overtime_minutes", "days_missing_in", "days_missing_out", "days_check",
           "days_absent", "days_off", "days_unscheduled", "total_hours")


def at(day: date, hms: str, days: int = 0) -> datetime:
    return datetime.combine(day + timedelta(days=days), datetime.strptime(hms, "%H:%M:%S").time())


def track() -> TrackState:
    t = TrackState(1)
    for i in range(3):
        t.add(Sample(0.9, Match(CODE, CODE, 0.8 + i / 100, None, None, "match"), np.zeros((8, 8, 3), np.uint8)))
    return t


def run_case(passes, first: date, last: date, roster=(), corrections=(), roster_after=()) -> dict:
    """The roster through the API, camera passes through the worker's Recorder (one per camera, like one worker
    process each), corrections and later roster changes through the API; returns what every screen, report and export
    shows."""
    conn = db.connect()
    db.sync_persons(conn, [CODE])
    conn.close()
    app = create_app()
    out = {}
    with TestClient(app):
        hr = as_role(app, "hr")
        set_roster(hr, roster)
        recorders, decisions = {}, []
        for cam, when in passes:
            rec = recorders.setdefault(cam, recognize.Recorder(CAMS[cam], 0, use_db=True))
            with contextlib.redirect_stdout(io.StringIO()):
                decisions.append(rec.record(when.strftime("%H:%M:%S"), when, track(), CODE, 3)[0])
        for rec in recorders.values():
            rec.conn.close()
        out["decisions"] = decisions
        for corr in corrections:
            body = dict(corr, reason="audit case")
            if "event_at" in body:
                wanted = body.pop("event_at")
                events = hr.get("/api/events", params={"from": first, "to": last + timedelta(days=1)}).json()["items"]
                body["event_id"] = next(e["id"] for e in events if e["timestamp"].endswith(wanted))
            assert hr.post("/api/attendance/corrections", json=body).status_code == 201
        set_roster(hr, roster_after)
        events = hr.get("/api/events", params={"from": first, "to": last + timedelta(days=1),
                                               "sort": "timestamp"}).json()["items"]
        out["events"] = [(e["timestamp"][11:19], e["direction"], e["repeat"], e["voided"], e["manual"]) for e in events]
        rows = hr.get("/api/attendance", params={"from": first, "to": last, "sort": "date"}).json()["items"]
        out["rows"] = [(r["date"], r["shift"], r["first_in"] and r["first_in"][11:19],
                        r["last_out"] and r["last_out"][11:19], r["worked_hours"], r["status"]) for r in rows]
        out["minutes"] = [(r["late_minutes"], r["early_leave_minutes"], r["overtime_minutes"]) for r in rows]
        out["overview"] = []
        day = first
        while day <= last:
            o = hr.get("/api/overview", params={"date": day}).json()
            out["overview"].append({k: o[k] for k in OVERVIEW})
            day += timedelta(days=1)
        daily = hr.get("/api/reports/daily", params={"from": first, "to": last}).json()["rows"]
        out["daily"] = [(r["date"], r["shift"], r["status"], r["worked_hours"]) for r in daily]
        people = hr.get("/api/reports/monthly", params={"from": first, "to": last}).json()["people"]
        out["monthly"] = {k: people[0][k] for k in MONTHLY} if people else None
        text = hr.get("/api/reports/daily/export", params={"from": first, "to": last, "format": "csv"})
        lines = list(csv.reader(io.StringIO(text.content.decode("utf-8-sig"))))
        assert (lines[0][4], lines[0][-1]) == ("Shift", "Status")
        out["csv"] = [(line[4], line[-1]) for line in lines[1:]]
        book = load_workbook(io.BytesIO(hr.get("/api/reports/monthly/export",
                                               params={"from": first, "to": last, "format": "xlsx"}).content))
        out["xlsx"] = [(row[4], row[-1]) for row in book["Daily detail"].iter_rows(min_row=2, values_only=True)]
    out["today_py"] = []
    day = first
    while day <= last:
        buf = io.StringIO()
        sys.argv = ["today.py", "--date", day.isoformat()]
        with contextlib.redirect_stdout(buf):
            today_script.main()
        lines = buf.getvalue().splitlines()
        header = next((line for line in lines if line.startswith("Name")), None)
        if header:
            out["today_py"] += [line[header.index("Status"):].strip() for line in lines if line.startswith(CODE)]
        day += timedelta(days=1)
    return out


def set_roster(client: TestClient, roster) -> None:
    """(date, shift name | "OFF" | None = remove) entries through PUT /api/roster."""
    if not roster:
        return
    shifts = {s["name"]: s["id"] for s in client.get("/api/shifts").json()["items"]}
    person = client.get("/api/people").json()["items"][0]["id"]
    entries = [{"person_id": person, "date": day.isoformat(), "shift_id": shifts.get(name), "clear": name is None}
               for day, name in roster]
    r = client.put("/api/roster", json={"entries": entries, "reason": "audit case roster"})
    assert r.status_code == 200, r.text


def same_everywhere(out, statuses):
    """The status of each day agrees across the API, the register, the CSV, the Excel sheet and today.py."""
    assert [r[5] for r in out["rows"]] == [d[2] for d in out["daily"]] == statuses
    assert [c[1] for c in out["csv"]] == [x[1] for x in out["xlsx"]] == out["today_py"] == statuses


def general(first=D, last=None):
    last = last or first
    return [(first + timedelta(days=i), "General") for i in range((last - first).days + 1)]


def test_a_in_and_out():
    out = run_case([("IN", at(D, "08:00:00")), ("OUT", at(D, "16:00:00"))], D, D, general())
    assert out["rows"] == [(str(D), "General", "08:00:00", "16:00:00", 8.0, "on time")]
    assert out["minutes"] == [(0, 0, 0)]  # 16:00 is end - grace: not early
    assert [e[2] for e in out["events"]] == [False, False]  # nothing is a repeat
    assert out["overview"] == [{"people": 1, "present": 1, "inside_now": None, "on_time": 1, "late": 0,
                                "early_leave": 0, "overtime": 0, "missing_in": 0, "missing_out": 0, "check": 0,
                                "absent": 0, "not_in_yet": 0, "off_day": 0, "unscheduled": 0}]
    assert out["daily"] == [(str(D), "General", "on time", 8.0)] and out["csv"] == [("General", "on time")]
    assert out["monthly"] == {"days": 1, "days_scheduled": 1, "days_present": 1, "days_late": 0, "late_minutes": 0,
                              "days_early_leave": 0, "early_leave_minutes": 0, "overtime_minutes": 0,
                              "days_missing_in": 0, "days_missing_out": 0, "days_check": 0, "days_absent": 0,
                              "days_off": 0, "days_unscheduled": 0, "total_hours": 8.0}
    same_everywhere(out, ["on time"])


def test_a2_late_early_leave_and_overtime():
    out = run_case([("IN", at(D, "09:00:00")), ("OUT", at(D, "18:00:00")),
                    ("IN", at(D, "08:10:00", days=1)), ("OUT", at(D, "14:00:00", days=1))],
                   D, D + timedelta(days=1), general(D, D + timedelta(days=1)))
    assert [r[5] for r in out["rows"]] == ["late", "early leave"]
    assert out["minutes"] == [(60, 0, 90), (0, 150, 0)]
    assert out["overview"][0]["late"] == 1 and out["overview"][0]["overtime"] == 1
    assert out["overview"][1]["early_leave"] == 1
    m = out["monthly"]
    assert (m["days_late"], m["late_minutes"], m["days_early_leave"], m["early_leave_minutes"],
            m["overtime_minutes"]) == (1, 60, 1, 150, 90)
    same_everywhere(out, ["late", "early leave"])


def test_b_back_in_after_the_last_out_is_a_repeat():
    out = run_case([("IN", at(D, "08:00:00")), ("OUT", at(D, "16:00:00")), ("IN", at(D, "17:00:00"))], D, D,
                   general())
    assert out["rows"] == [(str(D), "General", "08:00:00", "16:00:00", 8.0, "on time")]
    assert out["events"] == [("08:00:00", "IN", False, False, False), ("16:00:00", "OUT", False, False, False),
                             ("17:00:00", "IN", True, False, False)]
    assert out["monthly"]["total_hours"] == 8.0 and out["monthly"]["days_missing_out"] == 0
    same_everywhere(out, ["on time"])


def test_c_standing_in_front_of_the_camera_records_one_in():
    out = run_case([("IN", at(D, "08:00:00")), ("IN", at(D, "08:00:20"))], D, D, general())
    assert out["decisions"] == ["event", "cool-down"]
    assert [e[:2] for e in out["events"]] == [("08:00:00", "IN")]
    assert out["rows"] == [(str(D), "General", "08:00:00", None, None, "missing OUT")]
    assert out["overview"][0]["missing_out"] == 1 and out["overview"][0]["present"] == 1
    same_everywhere(out, ["missing OUT"])


def test_c2_a_second_in_camera_is_in_the_same_cool_down():
    out = run_case([("IN", at(D, "08:00:00")), ("IN2", at(D, "08:00:20"))], D, D, general())
    assert out["decisions"] == ["event", "cool-down"]  # another camera, same direction: one IN
    assert [e[:2] for e in out["events"]] == [("08:00:00", "IN")]


def test_c3_after_the_cool_down_a_second_in_is_kept_as_a_repeat():
    out = run_case([("IN", at(D, "08:00:00")), ("IN", at(D, "08:06:00"))], D, D, general())
    assert [(e[0], e[2]) for e in out["events"]] == [("08:00:00", False), ("08:06:00", True)]
    assert out["rows"] == [(str(D), "General", "08:00:00", None, None, "missing OUT")]


def test_d_lunch_break_is_inside_the_day():
    out = run_case([("IN", at(D, "08:00:00")), ("OUT", at(D, "12:00:00")), ("IN", at(D, "13:00:00")),
                    ("OUT", at(D, "16:00:00"))], D, D, general())
    assert out["rows"] == [(str(D), "General", "08:00:00", "16:00:00", 8.0, "on time")]
    assert [e[2] for e in out["events"]] == [False, True, True, False]  # the lunch OUT and IN are repeats
    same_everywhere(out, ["on time"])


def test_e_seen_only_by_the_exit_camera():
    out = run_case([("OUT", at(D, "16:00:00"))], D, D, general())
    assert out["rows"] == [(str(D), "General", None, "16:00:00", None, "missing IN")]
    assert out["overview"][0]["present"] == 0 and out["overview"][0]["missing_in"] == 1
    assert out["monthly"]["days_present"] == 0 and out["monthly"]["days_missing_in"] == 1
    same_everywhere(out, ["missing IN"])


def test_f_in_only_on_a_past_day():
    out = run_case([("IN", at(D, "08:00:00"))], D, D, general())
    assert out["rows"] == [(str(D), "General", "08:00:00", None, None, "missing OUT")]
    assert out["monthly"]["days_present"] == 1 and out["monthly"]["days_missing_out"] == 1
    same_everywhere(out, ["missing OUT"])


def test_f_in_only_today_is_still_at_work_everywhere():
    """18:00 today, the OUT window open until 20:30: on time and inside, not missing OUT yet."""
    out = run_case([("IN", at(TODAY, "08:00:00"))], TODAY, TODAY, general(TODAY))
    assert out["rows"] == [(str(TODAY), "General", "08:00:00", None, None, "on time")]
    assert out["overview"][0]["inside_now"] == 1 and out["overview"][0]["present"] == 1
    assert out["monthly"]["days_missing_out"] == 0
    same_everywhere(out, ["on time"])


def test_f2_not_in_yet_today():
    """Night today starts at 22:00 (IN window from 20:00): at 18:00 nobody is absent yet."""
    out = run_case([], TODAY, TODAY, [(TODAY, "Night")])
    assert out["rows"] == [(str(TODAY), "Night", None, None, None, "not in yet")]
    assert out["overview"][0]["not_in_yet"] == 1 and out["overview"][0]["absent"] == 0
    same_everywhere(out, ["not in yet"])


def test_g_night_shift_is_one_day_dated_the_day_it_starts():
    out = run_case([("IN", at(D, "21:55:00")), ("OUT", at(D, "06:05:00", days=1))], D, D + timedelta(days=1),
                   [(D, "Night")])
    assert out["rows"] == [(str(D), "Night", "21:55:00", "06:05:00", 8.17, "on time")]
    assert out["minutes"] == [(0, 0, 5)]
    assert [o["people"] for o in out["overview"]] == [1, 0]  # nothing on the second date: no roster, no loose pass
    assert out["monthly"]["days"] == 1 and out["monthly"]["total_hours"] == 8.17
    same_everywhere(out, ["on time"])


def test_h_voided_in_and_corrections():
    out = run_case([("IN", at(D, "08:00:00")), ("OUT", at(D, "16:00:00"))], D, D, general(), corrections=[
        {"action": "void", "event_at": "08:00:00"},
        {"action": "add", "person_id": 1, "direction": "IN", "timestamp": at(D, "08:30:00").isoformat()},
        {"action": "change", "event_at": "16:00:00", "timestamp": at(D, "16:30:00").isoformat()}])
    assert out["events"] == [("08:00:00", "IN", False, True, False), ("08:30:00", "IN", False, False, True),
                             ("16:30:00", "OUT", False, False, False)]  # a voided event is never a repeat
    assert out["rows"] == [(str(D), "General", "08:30:00", "16:30:00", 8.0, "on time")]  # 08:30 = start + grace
    same_everywhere(out, ["on time"])


def test_x_out_before_in_is_check():
    """An OUT in the IN window and an IN after the midpoint: passes, but none where they count."""
    out = run_case([("OUT", at(D, "07:00:00")), ("IN", at(D, "12:30:00"))], D, D, general())
    assert out["rows"] == [(str(D), "General", "12:30:00", "07:00:00", None, "check")]
    assert out["overview"][0]["check"] == 1 and out["overview"][0]["present"] == 0
    assert out["monthly"]["days_check"] == 1 and out["monthly"]["total_hours"] == 0.0
    same_everywhere(out, ["check"])


def test_off_day_and_unscheduled_never_count_as_absent():
    out = run_case([("IN", at(D, "09:00:00", days=1)), ("OUT", at(D, "15:00:00", days=1))],
                   D, D + timedelta(days=1), [(D, "OFF")])
    assert out["rows"] == [(str(D), None, None, None, None, "off day"),
                           (str(D + timedelta(days=1)), None, "09:00:00", "15:00:00", None, "unscheduled")]
    assert out["minutes"] == [(None, None, None), (None, None, None)]  # nothing guessed
    assert (out["overview"][0]["off_day"], out["overview"][1]["unscheduled"]) == (1, 1)
    assert sum(o["absent"] for o in out["overview"]) == 0
    m = out["monthly"]
    assert (m["days"], m["days_scheduled"], m["days_off"], m["days_unscheduled"], m["days_absent"]) == (2, 0, 1, 1, 0)
    assert out["csv"] == [("OFF", "off day"), ("", "unscheduled")]
    same_everywhere(out, ["off day", "unscheduled"])


def test_a_roster_change_recalculates_every_screen_from_the_raw_events():
    out = run_case([("IN", at(D, "08:05:00")), ("OUT", at(D, "16:35:00"))], D, D, [(D, "Morning")],
                   roster_after=[(D, "General")])
    assert out["rows"] == [(str(D), "General", "08:05:00", "16:35:00", 8.5, "on time")]  # Morning would be late
    assert out["minutes"] == [(0, 0, 5)]
    same_everywhere(out, ["on time"])


def test_cool_down_does_not_reset_at_midnight():
    out = run_case([("IN", at(D, "23:58:00")), ("IN", at(D, "00:02:00", days=1))], D, D + timedelta(days=1),
                   [(D, "Night")])
    assert out["decisions"] == ["event", "cool-down"]  # 4 minutes apart: one pass, whatever the date
    assert out["rows"] == [(str(D), "Night", "23:58:00", None, None, "missing OUT")]
    assert out["minutes"] == [(118, None, None)]


def test_correction_with_a_time_zone_is_refused_not_a_crash():
    conn = db.connect()
    db.sync_persons(conn, [CODE])
    db.upsert_camera(conn, "entrance", 0, "IN")
    conn.close()
    app = create_app()
    with TestClient(app):
        hr = as_role(app, "hr")
        for ts in ("2026-10-20T08:30:00Z", "2026-10-20T08:30:00+05:00"):
            r = hr.post("/api/attendance/corrections", json={"action": "add", "person_id": 1, "direction": "IN",
                                                             "timestamp": ts, "reason": "x"})
            assert r.status_code == 422 and "without a time zone" in r.text
        assert hr.get("/api/events", params={"from": D, "to": D}).json()["total"] == 0


def test_the_overview_inside_now_counts_people_back_in_after_their_last_out():
    """Case b today (IN, OUT, IN again), the OUT window still open: on time, and in the building now."""
    out = run_case([("IN", at(TODAY, "08:00:00")), ("OUT", at(TODAY, "16:00:00")), ("IN", at(TODAY, "17:00:00"))],
                   TODAY, TODAY, general(TODAY))
    assert out["rows"][0][5] == "on time" and out["overview"][0]["inside_now"] == 1
