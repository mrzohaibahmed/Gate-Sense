"""Reports on roster days (app/reports.py): the daily register, per-person totals, and the CSV/Excel files."""

import io
from datetime import date, datetime, timedelta

import openpyxl
import pytest

from app import db, reports


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "attendance.db")
    db.sync_persons(c, ["ali", "waqas"])
    ids = db.person_ids(c)
    cin = db.upsert_camera(c, "entrance", 0, "IN")
    cout = db.upsert_camera(c, "exit", 1, "OUT")
    general = next(s.id for s in db.list_shifts(c) if s.name == "General")  # 08:00-16:30
    days = [date(2026, 10, d) for d in range(1, 6)]
    db.set_roster(c, [(ids["waqas"], d, general) for d in days[:4]] + [(ids["waqas"], days[4], None)]
                  + [(ids["ali"], d, general) for d in days[1:3]], "plan", "hr1")
    # waqas: two days 09:00-17:00 (late 60, overtime 30), two absent, then OFF; ali: one day without an OUT, one absent
    for d in (1, 2):
        db.add_event(c, ids["waqas"], cin, datetime(2026, 10, d, 9, 0), 0.8, None)
        db.add_event(c, ids["waqas"], cout, datetime(2026, 10, d, 17, 0), 0.8, None)
    db.add_event(c, ids["ali"], cin, datetime(2026, 10, 2, 8, 10), 0.7, None)
    yield c
    c.close()


def test_daily_register(conn):
    rows = reports.range_rows(conn, date(2026, 10, 2), date(2026, 10, 2))
    assert [r["employee_code"] for r in rows] == ["ali", "waqas"]
    ali, waqas = rows
    assert (ali["first_in"], ali["last_out"], ali["status"], ali["late_minutes"]) == (
        datetime(2026, 10, 2, 8, 10), None, "missing OUT", 0)
    assert (waqas["worked_hours"], waqas["status"], waqas["late_minutes"], waqas["overtime_minutes"]) == (
        8.0, "late", 60, 30)
    assert waqas["shift"] == "General" and waqas["shift_start"] == datetime(2026, 10, 2, 8, 0)


def test_person_totals(conn):
    rows = reports.range_rows(conn, date(2026, 10, 1), date(2026, 10, 5))
    ali, waqas = reports.person_totals(rows)
    assert (waqas["days"], waqas["days_scheduled"], waqas["days_present"], waqas["days_late"],
            waqas["days_absent"], waqas["days_off"]) == (5, 4, 2, 2, 2, 1)
    assert (waqas["late_minutes"], waqas["overtime_minutes"], waqas["total_hours"], waqas["avg_hours"]) == (
        120, 60, 16.0, 8.0)
    assert (ali["days"], ali["days_present"], ali["days_missing_out"], ali["days_absent"], ali["avg_hours"]) == (
        2, 1, 1, 1, None)


def test_excel_and_csv_export(conn):
    rows = reports.range_rows(conn, date(2026, 10, 1), date(2026, 10, 5))
    totals = reports.person_totals(rows)
    wb = openpyxl.load_workbook(io.BytesIO(reports.rows_xlsx(
        [("Summary", reports.TOTAL_COLUMNS, totals), ("Daily detail", reports.DAILY_COLUMNS_RANGE, rows)], "t")))
    assert wb.sheetnames == ["Summary", "Daily detail"]
    assert wb["Summary"]["A1"].value == "Employee code" and wb["Summary"].freeze_panes == "A2"
    assert wb["Daily detail"].max_row == len(rows) + 1
    csv = reports.rows_csv(reports.DAILY_COLUMNS_RANGE, rows).decode("utf-8-sig").splitlines()
    assert csv[0] == ("Date,Employee code,Name,Department,Shift,Shift start,Shift end,First IN,Last OUT,Hours worked,"
                      "Late (min),Early leave (min),Overtime (min),Status")
    assert "2026-10-02,ali,ali,,General,08:00:00,16:30:00,08:10:00,,,0,,,missing OUT" in csv
    assert "2026-10-01,waqas,waqas,,General,08:00:00,16:30:00,09:00:00,17:00:00,8.0,60,0,30,late" in csv
    assert "2026-10-05,waqas,waqas,,OFF,,,,,,,,,off day" in csv


def test_a_night_shift_is_one_row_dated_the_day_it_starts(conn):
    ids = db.person_ids(conn)
    night = next(s.id for s in db.list_shifts(conn) if s.name == "Night")
    mon = date(2026, 10, 12)
    db.set_roster(conn, [(ids["ali"], mon, night)], "night week", "hr1")
    cin, cout = (r["id"] for r in db.list_cameras(conn))
    db.add_event(conn, ids["ali"], cin, datetime(2026, 10, 12, 21, 55), 0.8, None)
    db.add_event(conn, ids["ali"], cout, datetime(2026, 10, 13, 6, 5), 0.8, None)
    rows = reports.range_rows(conn, mon, mon + timedelta(days=1))
    assert [(r["date"], r["shift"], r["status"], r["worked_hours"]) for r in rows] == [
        (mon, "Night", "on time", 8.17)]
