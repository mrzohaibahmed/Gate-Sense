"""The API for the features the old Streamlit dashboard had and the React app lacked: live decisions over the WebSocket
(admins), person edit, void/restore, photo upload, delete photo / face data, camera delete/move/log, reports."""

import base64
import csv
import io
from datetime import date, datetime, timedelta

import cv2
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import auth
from api.live import decision_out
from api.main import create_app
from api.workers import WorkerManager
from app import camera_settings, config, db, enrollment
from app.workers import WorkerManager as AppWorkerManager
from tests.fakes import FakeDetector, FakeRecognizer, photo
from tests.test_api import FAKE, wait_for
from tests.test_data_api import (  # noqa: F401  (first_password, fixed_clock: fixtures)
    admin, as_role, at, clip, first_password, fixed_clock, items, seed,
)

TODAY, YESTERDAY = date.today(), date.today() - timedelta(days=1)


@pytest.fixture
def app():
    application = create_app(models=lambda: (FakeDetector(), FakeRecognizer()))
    with TestClient(application):
        application.state.seed = seed()
        yield application


def b64(img) -> str:
    return "data:image/jpeg;base64," + base64.b64encode(cv2.imencode(".jpg", img)[1].tobytes()).decode()


def audit(app, action: str) -> list[dict]:
    return [e for e in app.state.db.audit_entries() if e["action"] == action]


# --- 1. Live decisions (admins only) ------------------------------------------------------------------------


def test_decision_margin():
    out = decision_out({"time": "09:00:01", "kind": "match", "track": 3, "name": "waqas", "score": 0.712,
                        "second": 0.42, "face_px": 150, "detail": "vote 1/3", "liveness": 0.97})
    assert out["margin"] == 0.292 and out["face_px"] == 150 and out["liveness"] == 0.97
    spoof = decision_out({"kind": "skip", "detail": "spoof (0.12)", "liveness": 0.12, "face_px": 140})
    assert (spoof["detail"], spoof["liveness"]) == ("spoof (0.12)", 0.12)
    assert decision_out({"kind": "skip", "detail": "too small (60px < 80)"})["liveness"] is None
    assert decision_out({"kind": "skip", "detail": "too small (60px < 80)"})["margin"] is None


@pytest.fixture
def deciding_app(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "WS_POLL_SECONDS", 0.05)
    monkeypatch.setattr(config, "WS_CAMERA_SECONDS", 0.2)
    video = clip(tmp_path / "c.avi")
    camera_settings.save([{"name": "c", "direction": "OUT", "mirror": False, "type": "file", "path": str(video),
                           "source": str(video)}])
    run, logs, live_dir = tmp_path / "run", tmp_path / "logs", tmp_path / "live"
    inner = AppWorkerManager(None, script=FAKE, run_dir=run, log_dir=logs, live_dir=live_dir,
                             extra_args=("--no-window", "--run-dir", str(run), "--frames", "--decisions",
                                         "--live-dir", str(live_dir)))
    application = create_app(workers=WorkerManager(inner))
    try:
        with TestClient(application):
            yield application
    finally:
        for proc in inner._procs.values():
            proc.kill()


def messages(ws, n: int = 60) -> list[dict]:
    return [ws.receive_json() for _ in range(n)]


def test_decisions_reach_admins_only(deciding_app):
    a, hr = as_role(deciding_app, "admin"), as_role(deciding_app, "hr")
    cam = items(a.get("/api/cameras"))[0]
    assert a.post(f"/api/cameras/{cam['id']}/start").status_code == 200
    wait_for(lambda: a.get(f"/api/cameras/{cam['id']}").json()["worker"]["state"] == "running")
    with a.websocket_connect("/api/ws") as ws:
        for _ in range(300):
            m = ws.receive_json()
            if m["type"] == "decisions" and m["decisions"]:
                break
        assert m["camera"] == "c" and m["id"] == cam["id"]
        match = next(d for d in m["decisions"] if d["kind"] == "match")
        skip = next(d for d in m["decisions"] if d["kind"] == "skip")
        assert (match["name"], match["score"], match["second"], match["margin"], match["face_px"]) == (
            "waqas", 0.71, 0.42, 0.29, 150)
        assert (skip["detail"], skip["face_px"], skip["margin"]) == ("too small (60px < 80)", 60, None)
    with hr.websocket_connect("/api/ws") as ws:
        kinds = {m["type"] for m in messages(ws)}
        assert "camera" in kinds and "decisions" not in kinds  # HR has live:view but not live:decisions


# --- 2 and 3. Person edit, void / restore (existing endpoints) -----------------------------------------------


def test_person_edit_keeps_the_code_and_is_hr_only(app):
    pid = app.state.seed["people"]["1002"]
    hr = as_role(app, "hr")
    r = hr.patch(f"/api/people/{pid}", json={"name": "Bilal Khan", "department": "Accounts", "reason": "typo"})
    assert r.status_code == 200 and (r.json()["name"], r.json()["department"], r.json()["employee_code"]) == (
        "Bilal Khan", "Accounts", "1002")
    assert as_role(app, "viewer").patch(f"/api/people/{pid}", json={"name": "X", "reason": "x"}).status_code == 403
    assert audit(app, "edit_person")[0]["reason"] == "typo"


def test_void_and_restore(app):
    hr = as_role(app, "hr")
    eid = app.state.seed["events"]["b_in"]
    r = hr.post("/api/attendance/corrections", json={"action": "void", "event_id": eid, "reason": "wrong person"})
    assert r.status_code == 201 and r.json()["voided"] is True
    bilal = next(x for x in items(hr.get("/api/attendance")) if x["employee_code"] == "1002")
    assert bilal["status"] == "absent"  # a voided event doesn't count
    assert next(e for e in items(hr.get("/api/events")) if e["id"] == eid)["voided"] is True  # but stays listed
    assert hr.post("/api/attendance/corrections", json={"action": "restore", "event_id": eid,
                                                        "reason": "it was him"}).json()["voided"] is False
    assert as_role(app, "viewer").post("/api/attendance/corrections", json={
        "action": "void", "event_id": eid, "reason": "x"}).status_code == 403
    assert [e["action"] for e in audit(app, "void_event") + audit(app, "restore_event")] == ["void_event",
                                                                                            "restore_event"]


# --- 4. Upload photos ---------------------------------------------------------------------------------------


def test_upload_photos_saves_the_passing_ones(app):
    hr = as_role(app, "hr")
    pid = app.state.seed["people"]["1001"]
    body = {"reason": "photos from HR", "photos": [
        {"name": "front.jpg", "image": b64(photo("good", seed=1))},
        {"name": "group.jpg", "image": b64(photo("two"))},
        {"name": "blurry.jpg", "image": b64(photo("blurry"))},
        {"name": "side.jpg", "image": b64(photo("good", seed=2))},
        {"name": "broken.jpg", "image": "data:image/jpeg;base64,bm90IGFuIGltYWdl"},
    ]}
    r = hr.post(f"/api/people/{pid}/photos", json=body)
    assert r.status_code == 200, r.text
    res = {x["name"]: x for x in r.json()["results"]}
    assert [n for n, x in res.items() if x["ok"]] == ["front.jpg", "side.jpg"]
    assert res["front.jpg"]["saved_as"] == "1.jpg" and res["front.jpg"]["message"] == ""
    for bad in ("group.jpg", "blurry.jpg", "broken.jpg"):
        assert not res[bad]["ok"] and res[bad]["message"] and res[bad]["saved_as"] is None, bad
    assert r.json()["saved"] == 2 and r.json()["embeddings"] == 2
    assert sorted(p.name for p in enrollment.photo_paths(config.KNOWN_DIR / "1001")) == ["1.jpg", "2.jpg"]
    row = audit(app, "add_photos")[0]
    assert row["user"] == "hr" and row["reason"] == "photos from HR" and "group.jpg" in row["details"]
    assert hr.post(f"/api/people/{pid}/photos", json={**body, "photos": body["photos"] * 3}).status_code == 422
    assert as_role(app, "viewer").post(f"/api/people/{pid}/photos", json=body).status_code == 403


# --- 5. Delete a photo, delete face data ----------------------------------------------------------------------


def test_delete_one_photo_rebuilds(app):
    hr = as_role(app, "hr")
    pid = app.state.seed["people"]["1001"]
    hr.post(f"/api/people/{pid}/photos", json={"photos": [  # no reason: the default goes to the audit log
        {"name": f"{i}.jpg", "image": b64(photo("good", seed=i))} for i in range(3)]})
    assert enrollment.load_gallery().counts()["1001"] == 3
    assert hr.delete(f"/api/people/{pid}/photos/2.jpg").status_code == 422  # no reason
    r = hr.delete(f"/api/people/{pid}/photos/2.jpg", params={"reason": "blurred"})
    assert r.status_code == 200 and r.json()["photos"] == 2 and r.json()["embeddings"] == 2
    assert enrollment.load_gallery().counts()["1001"] == 2
    assert hr.delete(f"/api/people/{pid}/photos/9.jpg", params={"reason": "x"}).status_code == 404
    assert audit(app, "delete_photo")[0]["details"].count("2.jpg") == 1


def test_delete_face_data_keeps_the_person_and_attendance(app):
    hr = as_role(app, "hr")
    pid = app.state.seed["people"]["1001"]
    hr.post(f"/api/people/{pid}/photos", json={"reason": "x", "photos": [
        {"name": f"{i}.jpg", "image": b64(photo("good", seed=i))} for i in range(2)]})
    events_before = hr.get("/api/events", params={"person": pid, "from": YESTERDAY.isoformat()}).json()["total"]
    r = hr.post(f"/api/people/{pid}/delete-face-data", json={"reason": "left the company"})
    assert r.json() == {"photos": 2, "embeddings": 2, "deactivated": False}
    person = hr.get(f"/api/people/{pid}").json()
    assert (person["photos"], person["embeddings"], person["active"]) == (0, 0, True)
    assert hr.get("/api/events", params={"person": pid, "from": YESTERDAY.isoformat()}).json()["total"] == events_before
    assert not (config.KNOWN_DIR / "1001").exists() and "1001" not in enrollment.load_gallery().counts()
    assert audit(app, "delete_face_data")[0]["reason"] == "left the company"
    r = hr.post(f"/api/people/{pid}/delete-face-data", json={"reason": "again", "deactivate": True})
    assert r.json()["deactivated"] is True and hr.get(f"/api/people/{pid}").json()["active"] is False


# --- 6. Cameras: delete, move, log ----------------------------------------------------------------------------


def test_camera_delete_move_and_log(worker_app):  # noqa: F811
    a = as_role(worker_app, "admin")
    ids = {c["name"]: c["id"] for c in items(a.get("/api/cameras"))}
    assert [c["name"] for c in items(a.get("/api/cameras"))] == ["a", "c"]

    moved = a.post(f"/api/cameras/{ids['c']}/move", json={"step": -1})
    assert moved.status_code == 200 and [c["name"] for c in moved.json()] == ["c", "a"]
    assert [c["name"] for c in camera_settings.entries()] == ["c", "a"]
    assert audit(worker_app, "move_camera")[0]["details"] == '{"position": [1, 0]}'

    a.post(f"/api/cameras/{ids['a']}/start")
    wait_for(lambda: a.get(f"/api/cameras/{ids['a']}").json()["worker"]["state"] == "running")
    log_dir = worker_app.state.workers.log_path("a").parent
    (log_dir / "a.log").write_text("".join(f"line {i}\n" for i in range(800)) +
                                   "could not open rtsp://admin:S3cretPw@10.0.0.5:554/x\n", encoding="utf-8")
    log = a.get(f"/api/cameras/{ids['a']}/log").json()
    assert len(log["lines"]) == 500 and log["lines"][-1] == "could not open rtsp://admin:***@10.0.0.5:554/x"
    assert log["lines"][0] == "line 301" and log["size_bytes"] > 0  # 801 lines: the last 500
    assert len(a.get(f"/api/cameras/{ids['a']}/log", params={"lines": 10}).json()["lines"]) == 10
    file = a.get(f"/api/cameras/{ids['a']}/log.txt")
    assert file.headers["content-disposition"].startswith('attachment; filename="gatesense-a-')
    assert "S3cretPw" not in file.text and file.text.count("\n") == 801

    busy = a.delete(f"/api/cameras/{ids['a']}", params={"reason": "moved"})
    assert busy.status_code == 409 and "running" in busy.json()["detail"]
    assert a.post(f"/api/cameras/{ids['a']}/stop").json()["result"] == "stopped cleanly"
    assert a.delete(f"/api/cameras/{ids['a']}").status_code == 422  # no reason
    assert a.delete(f"/api/cameras/{ids['a']}", params={"reason": "moved to the back gate"}).status_code == 204
    assert [c["name"] for c in items(a.get("/api/cameras"))] == ["c"]
    last = a.delete(f"/api/cameras/{ids['c']}", params={"reason": "x"})
    assert last.status_code == 409 and "at least one camera" in last.json()["detail"]
    assert audit(worker_app, "delete_camera")[0]["reason"] == "moved to the back gate"
    hr = as_role(worker_app, "hr")
    for r in (hr.get(f"/api/cameras/{ids['c']}/log"), hr.post(f"/api/cameras/{ids['c']}/move", json={"step": 1}),
              hr.delete(f"/api/cameras/{ids['c']}", params={"reason": "x"})):
        assert r.status_code == 403


from tests.test_data_api import worker_app  # noqa: E402, F401  (fixture)


# --- 7. Reports ---------------------------------------------------------------------------------------------------


def test_daily_register(app):
    v = as_role(app, "viewer")
    r = v.get("/api/reports/daily").json()
    assert (r["first"], r["last"]) == (TODAY.isoformat(), TODAY.isoformat())
    rows = {x["name"]: x for x in r["rows"]}
    assert (rows["Ayesha"]["worked_hours"], rows["Ayesha"]["status"], rows["Ayesha"]["shift"]) == (
        9.08, "on time", "General")
    # Bilal came in late today and hasn't gone out, the OUT window still open: late, not "missing OUT" yet
    assert (rows["Bilal"]["status"], rows["Bilal"]["late_minutes"], rows["Bilal"]["open"]) == ("late", 80, True)
    assert rows["Bilal"]["last_out"] is None
    assert rows["Chand"]["status"] == "missing IN" and rows["Dawood"]["status"] == "absent"
    stores = v.get("/api/reports/daily", params={"department": "Stores"}).json()["rows"]
    assert [x["name"] for x in stores] == ["Ayesha", "Chand"]
    assert v.get("/api/reports/daily", params={"site_id": 99}).json()["rows"] == []


def test_monthly_totals_match_the_daily_rows(app):
    hr = as_role(app, "hr")
    params = {"from": (TODAY - timedelta(days=6)).isoformat(), "to": TODAY.isoformat()}
    daily = hr.get("/api/reports/daily", params=params).json()["rows"]
    people = {p["employee_code"]: p for p in hr.get("/api/reports/monthly", params=params).json()["people"]}
    assert len(daily) == 4 * 2 and set(people) == {"1001", "1002", "1003", "1004"}  # rostered yesterday and today
    counted = {"days_on_time": "on time", "days_late": "late", "days_early_leave": "early leave",
               "days_missing_in": "missing IN", "days_missing_out": "missing OUT", "days_check": "check",
               "days_absent": "absent", "days_not_in_yet": "not in yet", "days_off": "off day",
               "days_unscheduled": "unscheduled"}
    for code, total in people.items():
        mine = [d for d in daily if d["employee_code"] == code]
        hours = [d["worked_hours"] for d in mine if d["worked_hours"] is not None]
        assert total["days"] == len(mine) == 2 and total["days_scheduled"] == 2
        assert total["total_hours"] == round(sum(hours), 2)
        assert total["days_present"] == sum(d["status"] in ("on time", "late", "early leave", "missing OUT")
                                            for d in mine)
        for key, status in counted.items():
            assert total[key] == sum(d["status"] == status for d in mine), key
        for key in ("late_minutes", "early_leave_minutes", "overtime_minutes"):
            assert total[key] == sum(d[key] or 0 for d in mine), key
        assert total["avg_hours"] == (round(total["total_hours"] / len(hours), 2) if hours else None)
    assert (people["1001"]["days_present"], people["1001"]["total_hours"]) == (2, 17.08)  # 8.00 + 9.08
    assert people["1002"]["days_missing_out"] == 0  # today's shift is still running: late, at work
    assert people["1002"]["days_present"] == 1 and people["1004"]["days_absent"] == 2


def test_exports_and_their_audit(app):
    hr = as_role(app, "hr")
    params = {"from": YESTERDAY.isoformat(), "to": TODAY.isoformat()}
    r = hr.get("/api/reports/daily/export", params={**params, "format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"] == (f'attachment; filename="gatesense-daily-{YESTERDAY}_{TODAY}.csv"')
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["Date", "Employee code", "Name", "Department", "Shift", "Shift start", "Shift end", "First IN",
                       "Last OUT", "Hours worked", "Late (min)", "Early leave (min)", "Overtime (min)", "Status"]
    assert len(rows) == 1 + 8  # 4 people, 2 days
    bilal = next(row for row in rows[1:] if row[2] == "Bilal" and row[0] == str(TODAY))
    assert (bilal[4], bilal[7], bilal[8], bilal[10], bilal[13]) == ("General", "09:20:00", "", "80", "late")

    x = hr.get("/api/reports/monthly/export", params={**params, "format": "xlsx", "department": "Stores"})
    assert x.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "gatesense-monthly-" in x.headers["content-disposition"] and "-Stores.xlsx" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Daily detail"]
    summary, detail = wb["Summary"], wb["Daily detail"]
    assert summary.freeze_panes == "A2" and summary["A1"].font.bold and summary["A1"].fill.fgColor.rgb.endswith("0E7C66")
    assert [c.value for c in summary[1]][:6] == ["Employee code", "Name", "Department", "Days", "Scheduled",
                                                 "Present"]
    assert [r[1] for r in summary.iter_rows(min_row=2, values_only=True)] == ["Ayesha", "Chand"]
    hours_col = [c.value for c in summary[1]].index("Total hours") + 1
    assert summary.cell(row=2, column=hours_col).value == 17.08  # Ayesha's total hours, a number
    first_in = detail.cell(row=2, column=8)
    assert first_in.number_format == "hh:mm" and str(first_in.value) == "08:00:00"
    exported = [e for e in app.state.db.audit_entries() if e["action"] == "export_report"]
    assert [(e["user"], e["target"]) for e in exported] == [("hr", "report:monthly"), ("hr", "report:daily")]
    assert '"format": "xlsx"' in exported[0]["details"] and '"department": "Stores"' in exported[0]["details"]
    assert as_role(app, "guard").get("/api/reports/daily").status_code == 403
    assert hr.get("/api/reports/daily", params={"from": TODAY.isoformat(), "to": YESTERDAY.isoformat()}).status_code == 422


def test_enrollment_wizard_holds_the_webcam(worker_app):  # noqa: F811
    a, hr, viewer = as_role(worker_app, "admin"), as_role(worker_app, "hr"), as_role(worker_app, "viewer")
    ids = {c["name"]: c["id"] for c in items(a.get("/api/cameras"))}
    assert viewer.post("/api/enrollment/webcam").status_code == 403

    claim = hr.post("/api/enrollment/webcam")
    assert claim.status_code == 200
    body = claim.json()
    assert body["expires_seconds"] == config.ENROLL_CAMERA_CLAIM_SECONDS
    assert body["renew_seconds"] == config.ENROLL_CAMERA_RENEW_SECONDS < body["expires_seconds"]

    busy = a.post(f"/api/cameras/{ids['a']}/start")
    assert busy.status_code == 409 and "in use by the enrollment wizard (hr)" in busy.json()["detail"]
    tested = a.post("/api/cameras/test", json={"type": "webcam", "index": 0, "direction": "IN"})
    assert tested.status_code == 409 and "enrollment wizard" in tested.json()["detail"]
    assert a.post(f"/api/cameras/{ids['c']}/start").status_code == 200  # a video file is fine

    renewed = hr.post("/api/enrollment/webcam", json={"claim_id": body["claim_id"]})
    assert renewed.json()["claim_id"] == body["claim_id"]
    assert hr.delete(f"/api/enrollment/webcam/{body['claim_id']}").status_code == 204
    assert hr.delete(f"/api/enrollment/webcam/{body['claim_id']}").status_code == 204  # already gone: fine

    assert a.post(f"/api/cameras/{ids['a']}/start").status_code == 200
    wait_for(lambda: a.get(f"/api/cameras/{ids['a']}").json()["worker"]["state"] == "running")
    refused = hr.post("/api/enrollment/webcam")
    assert refused.status_code == 409
    assert refused.json()["detail"] == "The webcam is in use by the a camera worker; stop it on the Live cameras page first."
