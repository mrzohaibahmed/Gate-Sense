"""People > Import (app/people_import.py, api/routes/people_import.py): Excel and CSV files of employee code, name,
department (and optionally shift, site); a preview first; then all rows or none, every change audited. The app
fixture has 1001 Ayesha (Stores), 1002 Bilal (Office), 1003 Chand (Stores), 1004 Dawood (Gate)."""

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app import config, db
from app.people_import import ImportFileError, parse
from tests.test_data_api import app, as_role, first_password, items  # noqa: F401  (fixtures)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(rows) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def preview(client, data: bytes, filename: str = "staff.xlsx"):
    return client.post("/api/people/import/preview", params={"filename": filename}, content=data,
                       headers={"content-type": XLSX if filename.endswith(".xlsx") else "text/csv"})


def do_import(client, data: bytes, filename: str = "staff.xlsx", reason: str = "HR list of October"):
    return client.post("/api/people/import", params={"filename": filename, "reason": reason}, content=data,
                       headers={"content-type": XLSX if filename.endswith(".xlsx") else "text/csv"})


def person(client, code: str) -> dict:
    return next(p for p in items(client.get("/api/people", params={"page_size": 500})) if p["employee_code"] == code)


# --- Reading the file ------------------------------------------------------------------------------------------


def test_headings_are_found_by_name_whatever_the_order_and_extra_columns_are_ignored():
    data = workbook([
        [None],  # a blank row above the headings
        ["Dept", "Full Name", "Remarks", "Emp No"],
        ["Stores", "Ayesha", "ok", 1001],  # Excel keeps 1001 as a number
        [None, None, None, None],  # a blank row in between
        ["Office", "Bilal", "", 1002.0],
    ])
    rows, columns, ignored = parse(data, "staff.xlsx")
    assert columns == ["employee_code", "name", "department"] and ignored == ["Remarks"]
    assert [(r.line, r.employee_code, r.name, r.department) for r in rows] == [
        (3, "1001", "Ayesha", "Stores"), (5, "1002", "Bilal", "Office")]


@pytest.mark.parametrize("text, encoding", [
    ("Employee code;Name;Department\n2001;Zoë Malik;Stores\n", "utf-8-sig"),  # Excel's CSV in some countries
    ("employee number,name,department\r\n2001,Zoë Malik,Stores\r\n", "cp1252"),  # Windows' ANSI codepage
    ("Code\tName\tDepartment\n2001\tZoë Malik\tStores\n", "utf-8"),
])
def test_csv_files(text, encoding):
    rows, columns, _ = parse(text.encode(encoding), "staff.csv")
    assert [(r.employee_code, r.name, r.department) for r in rows] == [("2001", "Zoë Malik", "Stores")]


@pytest.mark.parametrize("data, filename, message", [
    (b"", "x.csv", "empty"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "old.xls", ".xls"),
    (b"PK\x03\x04not really a zip", "x.xlsx", "Excel workbook"),
    (b"Name,Department\nAyesha,Stores\n", "x.csv", "employee code column"),
    (b"Employee code,Name\n", "x.csv", "no people under it"),
])
def test_files_that_cant_be_read(data, filename, message):
    with pytest.raises(ImportFileError, match=message):
        parse(data, filename)


def test_too_many_rows(monkeypatch):
    monkeypatch.setattr(config, "IMPORT_MAX_ROWS", 2)
    with pytest.raises(ImportFileError, match="at most 2"):
        parse(b"Code,Name\n1,a\n2,b\n3,c\n", "x.csv")


# --- Preview ---------------------------------------------------------------------------------------------------


def test_preview_says_what_would_happen_and_changes_nothing(app):
    hr = as_role(app, "hr")
    data = workbook([
        ["Employee code", "Name", "Department"],
        [1001, "Ayesha", "Stores"],  # unchanged
        [1002, "Bilal Ahmed", "Accounts"],  # update: name and department
        [1003, "Chand", None],  # update: the department column is blank -> cleared
        [2001, "Zara", "Gate"],  # new
        [2002, None, "Gate"],  # error: no name
        ["a/b", "Bad Code", "Gate"],  # error: not a folder name
        [2001, "Zara again", "Gate"],  # error: the same code twice
    ])
    r = preview(hr, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["counts"] == {"new": 1, "update": 2, "unchanged": 1, "error": 3}
    assert body["columns"] == ["employee_code", "name", "department"]
    rows = {row["line"]: row for row in body["rows"]}
    assert rows[2]["action"] == "unchanged"
    assert rows[3]["action"] == "update" and rows[3]["changes"] == {"name": ["Bilal", "Bilal Ahmed"],
                                                                  "department": ["Office", "Accounts"]}
    assert rows[4]["changes"] == {"department": ["Stores", None]}
    assert rows[5]["action"] == "new" and rows[5]["employee_code"] == "2001"
    assert rows[6]["errors"] == ["no name"]
    assert "employee code must start with a letter or digit" in rows[7]["errors"][0]
    assert rows[8]["errors"] == ["the same employee code as line 5"]
    assert person(hr, "1002")["name"] == "Bilal"  # nothing changed
    assert all(e["action"] != "import_people" for e in app.state.db.audit_entries())


def test_columns_the_file_doesnt_have_are_left_alone(app):
    hr = as_role(app, "hr")
    body = preview(hr, b"Employee code,Name\n1002,Bilal\n1003,Chand\n", "names.csv").json()
    assert body["counts"]["unchanged"] == 2  # no department column: departments aren't compared


def test_sites_shift_and_capitals(app):
    hr = as_role(app, "hr")
    a = as_role(app, "admin")
    assert a.post("/api/sites", json={"name": "Yard", "reason": "second gate"}).status_code == 201
    conn = db.connect()
    db.add_person(conn, "EMP7", "Seven", None, None, "test", "test")
    conn.close()
    data = workbook([
        ["Employee code", "Name", "Shift", "Site"],
        [1004, "Dawood", "Night", "yard"],  # site by name, any capitals
        [3001, "Omar", "Day", None],  # blank site: the default site
        ["emp7", "Seven", None, None],  # differs only in capitals from EMP7
        [3002, "Noor", None, "Mars"],
    ])
    rows = {r["line"]: r for r in preview(hr, data).json()["rows"]}
    assert rows[2]["action"] == "update" and rows[2]["changes"] == {
        "shift": [None, "Night"], "site": [config.DEFAULT_SITE_NAME, "Yard"]}
    assert rows[3]["action"] == "new"
    assert "differs only in capitals from the existing employee code 'EMP7'" in rows[4]["errors"][0]
    assert rows[5]["errors"] == ["no site called 'Mars' (add it in Settings > Sites first)"]


# --- Import ----------------------------------------------------------------------------------------------------


def test_a_file_with_errors_imports_nothing(app):
    hr = as_role(app, "hr")
    r = do_import(hr, b"Employee code,Name,Department\n2001,Zara,Gate\n2002,,Gate\n", "staff.csv")
    assert r.status_code == 422 and "Nothing was imported" in r.json()["detail"] and "line 3: no name" in r.json()[
        "detail"]
    assert not any(p["employee_code"] == "2001" for p in items(hr.get("/api/people", params={"page_size": 500})))


def test_import_adds_updates_and_audits_everything_in_one_go(app):
    hr = as_role(app, "hr")
    conn = db.connect()
    pid = db.active_person_id(conn, "1004")
    db.set_person_active(conn, pid, False, "left the company", "admin")
    conn.close()
    data = workbook([
        ["Employee code", "Name", "Department"],
        [1001, "Ayesha", "Stores"],
        [1002, "Bilal Ahmed", "Accounts"],
        [1004, "Dawood", "Security"],  # inactive: updated, not reactivated
        [2001, "Zara", "Gate"],
        [2002, "Omar", None],
    ])
    rows = {r["line"]: r for r in preview(hr, data).json()["rows"]}
    assert rows[4]["notes"] == ["inactive: updated, but not reactivated"]
    r = do_import(hr, data)
    assert r.status_code == 200, r.text
    assert r.json() == {"file": "staff.xlsx", "added": 2, "updated": 2, "unchanged": 1}

    assert (person(hr, "1002")["name"], person(hr, "1002")["department"]) == ("Bilal Ahmed", "Accounts")
    zara = person(hr, "2001")
    assert (zara["name"], zara["department"], zara["active"], zara["site"]) == ("Zara", "Gate", True,
                                                                              config.DEFAULT_SITE_NAME)
    assert zara["photos"] == 0 and zara["embeddings"] == 0  # faces still to enroll
    dawood = person(hr, "1004")
    assert dawood["department"] == "Security" and dawood["active"] is False
    assert person(hr, "1003")["department"] == "Stores"  # not in the file: left alone

    audit = [e for e in app.state.db.audit_entries() if e["user"] == "hr" and e["reason"] == "HR list of October"]
    assert sorted(e["action"] for e in audit) == ["add_person", "add_person", "edit_person", "edit_person",
                                                  "import_people"]
    summary = next(e for e in audit if e["action"] == "import_people")
    assert summary["target"] == "file:staff.xlsx"
    assert json.loads(summary["details"]) == {"added": 2, "updated": 2, "unchanged": 1, "rows": 5}
    edit = next(e for e in audit if e["target"] == "person:1002")
    assert json.loads(edit["details"])["department"] == ["Office", "Accounts"]

    again = do_import(hr, data)  # the same file twice: nothing left to do
    assert again.json() == {"file": "staff.xlsx", "added": 0, "updated": 0, "unchanged": 5}


def test_import_needs_a_reason_and_people_manage(app):
    hr, viewer = as_role(app, "hr"), as_role(app, "viewer")
    data = b"Employee code,Name\n2001,Zara\n"
    assert do_import(hr, data, "x.csv", reason="   ").status_code == 422
    assert hr.post("/api/people/import", params={"filename": "x.csv"}, content=data).status_code == 422
    assert preview(viewer, data, "x.csv").status_code == 403
    assert do_import(viewer, data, "x.csv").status_code == 403
    assert viewer.get("/api/people/import/template").status_code == 403


def test_unreadable_file_is_a_clear_422(app):
    r = preview(as_role(app, "hr"), b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "old.xls")
    assert r.status_code == 422 and "save it as .xlsx" in r.json()["detail"]


# --- Template ----------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("kind", ["xlsx", "csv"])
def test_the_template_is_a_file_the_import_reads(app, kind):
    hr = as_role(app, "hr")
    r = hr.get("/api/people/import/template", params={"format": kind})
    assert r.status_code == 200 and f'filename="gatesense-people.{kind}"' in r.headers["content-disposition"]
    body = preview(hr, r.content, f"gatesense-people.{kind}").json()
    assert body["columns"] == ["employee_code", "name", "department", "shift", "site"]
    [row] = body["rows"]
    assert (row["employee_code"], row["action"]) == ("1001", "update")  # the example is 1001 Ayesha Khan
    if kind == "xlsx":
        assert load_workbook(io.BytesIO(r.content)).active["A2"].number_format == "@"  # codes as text


def test_import_routes_dont_clash_with_a_person_id(app):
    hr = as_role(app, "hr")
    assert hr.get("/api/people/import/template").status_code == 200  # not "/api/people/{person_id}"
    assert hr.get(f"/api/people/{app.state.seed['people']['1001']}").status_code == 200
